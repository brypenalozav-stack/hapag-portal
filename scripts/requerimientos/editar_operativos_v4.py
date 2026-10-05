"""Fase 2: documentos operativos v4 (Pendientes, NexusV2 y Comparación).

Lee los originales de katu (nunca los escribe) y genera en katu\\v4\\:
- Portal_2_0_Pendientes_v4.xlsx: responsables según Q5, columnas nuevas, reformulaciones, tareas
  nuevas desde el ID 119, fechas solo para TAREAS_PLAN (config.CALENDARIO) y hoja Resumen con el
  gráfico recreado.
- Funcionalidades NexusV2_v4.xlsx: RACI propuesto para las filas vacías, Q3 en "Datos" y columnas nuevas.
- Comparacion_Requerimientos_Por_Fase_v4.xlsx: compara contra el código de hapag-portal (línea base),
  sin las hojas del sitio externo.

No hay Gantt_v4: el Gantt de macros queda fuera del alcance del plan.
"""
import copy
import datetime
import json
import os
import re
import sys

import docx
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import PatternFill
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import TableColumn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402
import docx_utils  # noqa: E402
from linea_base import DOCS, REPO  # noqa: E402

ORIG_PENDIENTES = os.path.join(config.KATU, "Portal_2_0_Pendientes.xlsx")
ORIG_NEXUS = os.path.join(config.KATU, "Funcionalidades NexusV2.xlsx")
ORIG_COMPARACION = os.path.join(config.KATU, "Comparacion_POC_Requerimientos_Por_Fase.xlsx")
ESPEC_V4 = os.path.join(config.V4, "Portal_2.0_Especificacion_Funcional_v4.docx")
MATRIZ_BASE = os.path.join(config.V4, "Matriz_Fichas_Linea_Base_v4.xlsx")
INTEGRACIONES = os.path.join(REPO, "docs", "integraciones")

HOJA_PENDIENTES = "Pendientes Portal 2.0"
SIN_FECHA = "Sin fecha en fuentes"
FUERA_ALCANCE_MACROS = "Fuera del alcance del desarrollo Portal 2.0 (área Macros/RPX)"
ORIGENES = ["RACI NexusV2", "Plan Portal 2.0", "Propuesta", "Propuesta – sin nombre en fuentes"]
FORMATO_FECHA = "dd-mm-yyyy"

# ---------------------------------------------------------------- personas y áreas (Q5)

AREA_PERSONA = {
    "Fer": "Finanzas", "Ricardo": "Finanzas",
    "Kari": "Comercial",
    "Cami/Mati": "Customer Service", "Cami": "Customer Service", "Mati": "Customer Service",
    "Lucho": "Nexus/IT", "Jorge": "Nexus/IT",
    "Andrés": "Arquitectura/QA", "Andres": "Arquitectura/QA",
    "Diego": "Macros/RPX", "RPX": "Macros/RPX",
    "Katu": "Producto/Negocio",
}
LEGAL = "LEGAL"
SEGURIDAD = "SEGURIDAD"
SIN_NOMBRE = {LEGAL: ("Área Legal", "Área Legal (titular del área)"),
              SEGURIDAD: ("Área Seguridad TI", "Área Seguridad TI (titular del área)")}
EQUIPO = ("Desarrollo", "Desarrollo – Equipo de desarrollo Portal 2.0")


def area_de(nombre):
    if nombre in SIN_NOMBRE:
        return SIN_NOMBRE[nombre][0]
    return AREA_PERSONA[nombre]


def personas(nombres):
    """'Área – Nombre, Nombre + Área – Nombre' agrupando por área y conservando el orden."""
    grupos = []
    for n in nombres:
        if n in SIN_NOMBRE:
            grupos.append((SIN_NOMBRE[n][1], None))
            continue
        nombre = "Andrés" if n == "Andres" else n
        area = AREA_PERSONA[n]
        for g in grupos:
            if g[0] == area and g[1] is not None:
                g[1].append(nombre)
                break
        else:
            grupos.append((area, [nombre]))
    return " + ".join(a if ns is None else f"{a} – {', '.join(ns)}" for a, ns in grupos)


def celda_responsable(r, apoyo=None):
    """Formato Q5: `Área – Nombre (R); apoyo: Área – Nombre`. `r` y `apoyo` son listas de nombres o textos ya formateados."""
    def fmt(x):
        if isinstance(x, str):
            return x
        return personas(x)
    texto = f"{fmt(r)} (R)"
    if apoyo:
        texto += f"; apoyo: {fmt(apoyo)}"
    return texto


def origen_propuesta(r):
    primero = r[0] if isinstance(r, list) else None
    return ORIGENES[3] if primero in SIN_NOMBRE else ORIGENES[2]


def nombres_validadores(texto):
    """'Andrés, Kari (glosario EN)' -> ['Andrés', 'Kari (glosario EN)'] (nombres de config.CALENDARIO)."""
    partes = re.split(r",\s*|\s+y\s+", texto)
    return [p.strip() for p in partes if p.strip()]


def apoyo_fases(fases):
    vistos, salida = set(), []
    for f in fases:
        for v in nombres_validadores(config.CALENDARIO[f][2]):
            base = re.sub(r"\s*\(.*\)$", "", v)
            if base in vistos:
                continue
            vistos.add(base)
            extra = v[len(base):]
            salida.append(f"{AREA_PERSONA[base]} – {base}{extra}")
    return " + ".join(salida)


# ---------------------------------------------------------------- contratos (Fase 6a)

# ID, archivo relativo a docs/integraciones, sistema
CONTRATOS = [
    ("CT-NEXUS", "contratos/nexus.openapi.yaml", "Nexus"),
    ("CT-FIS", "contratos/fis.openapi.yaml", "FIS/Data Lake"),
    ("CT-KHIPU", "contratos/khipu.openapi.yaml", "Khipu"),
    ("CT-BCH", "contratos/banco-chile.openapi.yaml", "Banco de Chile"),
    ("CT-SANT", "contratos/santander.openapi.yaml", "Santander"),
    ("CT-BCI", "contratos/bci.openapi.yaml", "BCI"),
    ("CT-DBNET", "contratos/dbnet.openapi.yaml", "DBNet/SII"),
    ("CT-TRACK", "contratos/tracking.openapi.yaml", "Tracking"),
    ("CT-SIGN", "contratos/firma.openapi.yaml", "Firma"),
    ("CT-STORAGE", "storage.md", "Storage"),
    ("CT-MERC", "contratos/mercurio.openapi.yaml", "Mercurio"),
    ("CT-TATC", "contratos/tatc.openapi.yaml", "TATC/Flagare"),
    ("CT-NAVE", "navesoft.md", "Navesoft"),
    ("CT-DEP", "deposito.md", "Depósito"),
    ("CT-DISP", "dispute.md", "Dispute"),
]

# Sistemas con puerto (Integrations:<Sistema>:Mode) -> contrato
SISTEMAS_MODO = [
    ("Nexus", "CT-NEXUS", True), ("Fis", "CT-FIS", True), ("Khipu", "CT-KHIPU", True),
    ("BancoChile", "CT-BCH", True), ("Santander", "CT-SANT", False), ("Bci", "CT-BCI", False),
    ("DbNet", "CT-DBNET", True), ("Tracking", "CT-TRACK", True), ("Signature", "CT-SIGN", False),
    ("Storage", "CT-STORAGE", False),
]


def responsable_contrato(archivo):
    """Lee `info.x-responsable` (YAML) o la fila 'Responsable' (Markdown) del contrato."""
    ruta = os.path.join(INTEGRACIONES, archivo)
    with open(ruta, encoding="utf-8") as f:
        texto = f.read()
    if archivo.endswith(".yaml"):
        m = re.search(r"^\s*x-responsable:\s*'?(.*?)'?\s*$", texto, re.M)
    else:
        m = re.search(r"^\|\s*Responsable\s*\|\s*(.*?)\s*\|\s*$", texto, re.M)
    if not m:
        raise SystemExit(f"Sin responsable en {ruta}")
    return m.group(1)


def celda_desde_contrato(resp):
    """'Nexus/IT – Lucho (apoyo: Jorge)' -> (celda Q5, área, origen)."""
    m = re.match(r"^(.*?)(?:\s*\(apoyo:\s*(.*)\))?$", resp)
    r, apoyo = m.group(1).strip(), m.group(2)
    if r.startswith("Área Legal"):
        r_txt, area, origen = SIN_NOMBRE[LEGAL][1], SIN_NOMBRE[LEGAL][0], ORIGENES[3]
    elif r.startswith("Área Seguridad TI"):
        r_txt, area, origen = SIN_NOMBRE[SEGURIDAD][1], SIN_NOMBRE[SEGURIDAD][0], ORIGENES[3]
    else:
        r_txt, area, origen = r, r.split(" – ")[0], ORIGENES[2]
    if apoyo:
        apoyo = apoyo if " – " in apoyo else personas([a.strip() for a in apoyo.split(",")])
    else:
        apoyo = EQUIPO[1]
    return celda_responsable(r_txt, apoyo), area, origen


# ---------------------------------------------------------------- fichas, puertos y contratos

PUERTO = {
    "EXE": ("IExemptionReader", "DummyExemptionReader", "HttpNexusClient", "CT-NEXUS"),
    "CRED": ("ICreditConditionReader", "DummyCreditConditionReader", "HttpNexusClient", "CT-NEXUS"),
    "TC": ("IExchangeRateProvider", "DummyExchangeRateProvider", "HttpNexusClient", "CT-NEXUS"),
    "TAR": ("ITariffProvider", "DummyTariffProvider", "HttpNexusClient", "CT-NEXUS"),
    "SHIP": ("IShipmentSource", "DummyShipmentSource", "HttpShipmentSource", "CT-FIS"),
    "PAY": ("IPaymentProvider", "DummyPaymentProvider", "HttpKhipuPaymentProvider, HttpBancoChilePaymentProvider",
            "CT-KHIPU, CT-BCH, CT-SANT, CT-BCI"),
    "INV": ("IInvoiceProvider", "DummyInvoiceProvider", "HttpInvoiceProvider", "CT-DBNET"),
    "TRK": ("ITrackingProvider", "DummyTrackingProvider", "HttpTrackingProvider", "CT-TRACK"),
    "SIGN": ("IDocumentSigner", "DummyDocumentSigner", None, "CT-SIGN"),
    "STO": ("IFileStorage", "DummyFileStorage", None, "CT-STORAGE"),
}
PUERTOS_FICHA = {
    "M4-01": ["EXE"], "M4-02": ["EXE"],
    "M4-03": ["CRED"], "M5-07": ["CRED"], "M8-02": ["CRED"], "M8-03": ["CRED"],
    "M5-04": ["TC"], "M5-05": ["TC"], "M8-01": ["TAR"],
    "M2-01": ["SHIP"], "M2-02": ["SHIP"], "M2-06": ["SHIP"], "M2-07": ["SHIP"], "M2-09": ["SHIP"],
    "M3-03": ["SHIP"], "M3-18": ["SHIP"],
    "M5-01": ["PAY"], "M5-03": ["PAY"], "M7-02": ["PAY", "INV"], "NF-01": ["PAY"], "NF-04": ["PAY"],
    "NF-12": ["PAY"],
    "M3-11": ["INV"], "M5-09": ["INV"], "M7-01": ["INV"],
    "M2-08": ["TRK"],
    "M6-01": ["SIGN"], "M6-02": ["SIGN"], "M6-07": ["SIGN"],
    "M1-07": ["STO"], "M5-06": ["STO"], "M6-09": ["STO"], "NF-16": ["STO"],
}
# Código de las fases 6b/6c que no es un puerto de integración.
CODIGO_EXTRA = {
    "NF-07": "HTTPS/HSTS y ForwardedHeaders en Program.cs (6c)",
    "NF-09": "SecretTypes con claves de integración (6b)",
    "NF-11": "Resiliencia estándar de HttpClient en DI.Integrations.Partial (6c); errores Integration.* (6b)",
    "NF-27": "IntegrationLoggingHandler con X-Correlation-Id y métrica hapagportal.integrations.errors (6c)",
    "NF-12": "Resiliencia estándar en clientes de pago Real (6c)",
}
CONTRATOS_EXTRA = {
    "M2-05": "CT-DISP", "M2-09": "CT-FIS, CT-TATC", "M3-04": "CT-MERC, CT-TATC", "M3-05": "CT-MERC, CT-TATC",
    "M3-06": "CT-MERC, CT-TATC", "M3-16": "CT-TATC", "M6-07": "CT-SIGN, CT-TATC", "M5-06": "CT-DEP, CT-STORAGE",
    "NF-11": "CT-NEXUS, CT-FIS", "NF-09": "Todos los contratos con clave de API",
}


def contrato_ficha(fid):
    if fid in CONTRATOS_EXTRA:
        return CONTRATOS_EXTRA[fid]
    ids = []
    for p in PUERTOS_FICHA.get(fid, []):
        for c in PUERTO[p][3].split(", "):
            if c not in ids:
                ids.append(c)
    return ", ".join(ids) if ids else "—"


def codigo_puertos(fid):
    partes = []
    for p in PUERTOS_FICHA.get(fid, []):
        iface, dummy, real, _ = PUERTO[p]
        adaptadores = dummy + (f", {real}" if real else "")
        partes.append(f"{iface} ({adaptadores})")
    texto = ""
    if partes:
        texto = "Fases 6b/6c: " + "; ".join(partes)
    if fid in CODIGO_EXTRA:
        texto = (texto + "; " if texto else "Fases 6b/6c: ") + CODIGO_EXTRA[fid]
    return texto


def fichas_v4():
    d = docx.Document(ESPEC_V4)
    out = docx_utils.fichas(d)
    if len(out) != 134:
        raise SystemExit(f"La especificación v4 tiene {len(out)} fichas (se esperaban 134)")
    return out


def matriz_base():
    """{id: (cobertura, evidencia)} desde Matriz_Fichas_Linea_Base_v4.xlsx (Fase 0)."""
    ws = openpyxl.load_workbook(MATRIZ_BASE, read_only=True).active
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    return {f[0]: (f[5], f[6]) for f in filas if f[0]}


def fase_de(encabezado):
    m = re.search(r"FASE\s+(\d)", encabezado)
    return m.group(1) if m else "?"


def commit_base():
    with open(os.path.join(DOCS, "linea-base-katu.json"), encoding="utf-8") as f:
        datos = json.load(f)
    fecha = datetime.date.fromisoformat(datos["generado"][:10]).strftime("%d-%m-%Y")
    return datos["commit_repo"][:7], fecha


# ---------------------------------------------------------------- utilidades de hoja

def copiar_estilo(origen, destino):
    destino._style = copy.copy(origen._style)


def rehacer_merge(ws, viejo, nuevo):
    if viejo in [str(r) for r in ws.merged_cells.ranges]:
        ws.unmerge_cells(viejo)
    ws.merge_cells(nuevo)


def ampliar_cf(ws, ultima):
    nuevo = ConditionalFormattingList()
    for cf in ws.conditional_formatting:
        col = re.match(r"[A-Z]+", str(cf.sqref)).group(0)
        for regla in cf.rules:
            nuevo.add(f"{col}5:{col}{ultima}", regla)
    ws.conditional_formatting = nuevo


def ampliar_dv(ws, ultima):
    for dv in ws.data_validations.dataValidation:
        col = re.match(r"[A-Z]+", str(dv.sqref)).group(0)
        dv.sqref = MultiCellRange(f"{col}5:{col}{ultima}")


# ================================================================ NexusV2

RACI_PROPUESTO = {
    # función (prefijo de la columna A) -> (R, A); S Jorge, C Andres, I Katu / Kari
    "Datos": ("Lucho", "Jorge"),
    "Cuenta Admin": ("Lucho", "Jorge"),
    "Cuenta CS": ("Cami/Mati", "Kari"),
    "Errores de facturación": ("Lucho", "Fer"),
    "Plazos documentales": ("Lucho", "Cami/Mati"),
}

# función -> (criticidad Fase 1, fichas v4, estado en hapag-portal, contrato)
NEXUS_COLUMNAS = {
    "Creditos": ("Alta", "M4-03, M5-07, M8-02",
                 "P – backend/src/HapagPortal.Domain/Entities/CreditClient.cs:5-15 (CRUD local, sin lectura Nexus); "
                 "puerto ICreditConditionReader con DummyCreditConditionReader y HttpNexusClient (6b/6c)", "CT-NEXUS"),
    "Exenciones": ("Alta", "M4-01, M4-02",
                   "P – backend/src/HapagPortal.Domain/Entities/DemurrageExemption.cs:5-11 (CRUD local, sin consumidor); "
                   "puerto IExemptionReader con DummyExemptionReader y HttpNexusClient (6b/6c)", "CT-NEXUS"),
    "Confirmacion Pago": ("Alta", "M5-06, NF-04",
                          "P – backend/src/HapagPortal.Application/Payments/Commands/Webhooks/KhipuWebhookCommandHandler.cs:40-54 "
                          "(webhooks Khipu y Banco de Chile); sin confirmación de depósito desde Nexus",
                          "CT-KHIPU, CT-BCH, CT-SANT, CT-BCI, CT-DEP"),
    "Cambio de Almacen": ("Alta", "M3-04, M3-05, M3-06",
                          "P – backend/src/HapagPortal.Application/WarehouseChanges/Commands/Create/CreateWarehouseChangeCommandHandler.cs:21-55 "
                          "(solicitud individual con monto, sin adjunto)", "CT-MERC, CT-TATC, CT-NAVE"),
    "Generar TATC": ("Media", "M2-09", "N – sin código de TATC (GetBLByNumberQueryHandler.cs:26 consulta solo el BL)",
                     "CT-TATC"),
    "Generar CLD": ("Media", "M6-07, M3-16", "N – sin código de CLD", "CT-TATC"),
    "Conexion Portal": ("Alta", "Cap. 15 (integración)",
                        "N – sin conexión con Nexus; puertos Nexus con Dummy y cliente Real contra simulador "
                        "(backend/src/HapagPortal.Infrastructure/Integrations/Nexus/HttpNexusClient.cs:13)", "CT-NEXUS, CT-NAVE"),
    "API Estado de Cuenta": ("Baja (ficha Fase 2; M5-07 y M6-07 dependen de M4-03/M8-02, Q1)", "M7-03",
                             "N – sin estado de cuenta", "CT-NEXUS (operación por definir)"),
    "Datos": ("Alta", "Cap. 15 (integración, Q3)",
              "N – decidido por Q3: consulta bajo demanda vía API con caché corta; puertos Nexus de las fases 6b/6c",
              "CT-NEXUS"),
    "Cuenta Admin": ("Media", "M8-06",
                     "P – backend/src/HapagPortal.Infrastructure/Authentication/PermissionResolver.cs:10-31 "
                     "(Admin con todos los permisos)", "Sin contrato"),
    "Cuenta CS": ("Baja (ficha Fase 2)", "M8-05",
                  "P – backend/src/HapagPortal.Domain/Constants/RoleCodes.cs:9-17 (roles internos Coordinador/Supervisor)",
                  "Sin contrato"),
    "TC": ("Alta", "M5-05",
           "N – backend/src/HapagPortal.Domain/Entities/Currency.cs:10 (tasa sembrada, sin lectura); "
           "puerto IExchangeRateProvider con DummyExchangeRateProvider y HttpNexusClient (6b/6c)", "CT-NEXUS"),
    "Tarifa": ("Alta", "M8-01",
               "N – sin mantenedor; puerto ITariffProvider con DummyTariffProvider y HttpNexusClient (6b/6c)", "CT-NEXUS"),
    "Counter": ("Baja (ficha Fase 2)", "M8-09", "N – sin código de canje ni desconsolidado", "Sin contrato"),
    "Errores de facturación": ("No aplica (fuera del alcance, Q7)", "Sin ficha (Q7: permanece en Nexus)",
                               "N – fuera del alcance del portal", "Sin contrato (Q7)"),
    "Plazos documentales": ("Baja (ficha Fase 2)", "M2-10",
                            "N – backend/src/HapagPortal.Domain/Validation/DeadlineCalculator.cs:9-36 calcula plazos de Aduana, "
                            "no plazos documentales por nave", "CT-NEXUS (operación por definir)"),
}

NEXUS_Q3 = {
    "B": "Consultar a Nexus vía API bajo demanda, con caché corta en el portal (Q3). Nexus no escribe tablas del portal.",
    "C": "Decisión Q3 del registro de decisiones v4: el portal pide a Nexus caso por caso vía API; descartada la "
         "actualización de tablas del portal desde Nexus.",
    "M": "Resuelta por Q3: API bajo demanda con caché corta, sin acceso directo a BD.",
}

NEXUS_CABECERAS_NUEVAS = ["Criticidad Fase 1", "Ficha spec v4", "Estado en hapag-portal (C/P/N, file:line)",
                          "Contrato de integración", "Origen RACI"]


def clave_funcion(nombre, claves):
    for k in claves:
        if nombre.startswith(k):
            return k
    return None


def raci_v4():
    """[(fila, función, R, A, S, origen)] de NexusV2 con el RACI propuesto en las filas vacías."""
    ws = openpyxl.load_workbook(ORIG_NEXUS).active
    filas = []
    for r in range(3, ws.max_row + 1):
        nombre = ws.cell(r, 1).value
        if not nombre:
            continue
        R, A, S = ws.cell(r, 5).value, ws.cell(r, 6).value, ws.cell(r, 7).value
        clave = clave_funcion(nombre, RACI_PROPUESTO)
        if not R and clave:
            R, A = RACI_PROPUESTO[clave]
            S = "Jorge"
            origen = "Propuesta"
        else:
            origen = "RACI NexusV2"
        filas.append((r, nombre, R, A, S, origen))
    return filas


def nexus():
    wb = openpyxl.load_workbook(ORIG_NEXUS)
    ws = wb.active
    ultima_col = ws.max_column  # M
    for i, cab in enumerate(NEXUS_CABECERAS_NUEVAS, start=ultima_col + 1):
        c = ws.cell(2, i, cab)
        copiar_estilo(ws.cell(2, ultima_col), c)
    for fila, nombre, R, A, S, origen in raci_v4():
        if origen == "Propuesta":
            ws.cell(fila, 5, R)
            ws.cell(fila, 6, A)
            ws.cell(fila, 7, S)
            ws.cell(fila, 8, "Andres")
            ws.cell(fila, 9, "Katu / Kari")
        if nombre == "Datos":
            for col, texto in NEXUS_Q3.items():
                ws[f"{col}{fila}"] = texto
        clave = clave_funcion(nombre, NEXUS_COLUMNAS)
        if clave is None:
            raise SystemExit(f"NexusV2: función sin columnas v4: {nombre}")
        valores = list(NEXUS_COLUMNAS[clave]) + [origen]
        for i, v in enumerate(valores, start=ultima_col + 1):
            c = ws.cell(fila, i, v)
            copiar_estilo(ws.cell(fila, ultima_col), c)
    for col, ancho in zip("NOPQR", [20, 24, 60, 28, 16]):
        ws.column_dimensions[col].width = ancho
    salida = config.ruta_v4("Funcionalidades NexusV2_v4.xlsx")
    wb.save(salida)
    return salida


# ================================================================ Pendientes

MODULOS_NUEVOS = {
    range(1, 5): "Transición y línea base",
    range(9, 13): "Herramientas de desarrollo heredadas",
    range(19, 25): "Base de datos (PostgreSQL)",
}

TEXTOS = {
    9: "Inventariar y migrar herramientas de desarrollo heredadas",
    19: "Revisar el funcionamiento de la base de datos PostgreSQL (EF Core con Npgsql)",
    20: "Validar la configuración de PostgreSQL y las migraciones de HapagPortal.DatabaseMigrations para Portal 2.0",
    37: "Datos – Definir el modelo de consulta bajo demanda Portal→Nexus vía API (con caché), sin réplica de tablas",
    42: "Validar la API FIS de organizaciones (consulta bajo demanda vía IShipmentSource; entrada transitoria por "
        "POST bills-of-lading/import)",
    43: "Validar la API FIS de datos (consulta bajo demanda vía IShipmentSource; entrada transitoria por "
        "POST bills-of-lading/import)",
    44: "Validar la API FIS de documentos (consulta bajo demanda vía IShipmentSource; entrada transitoria por "
        "POST bills-of-lading/import)",
}

DEPENDENCIAS = {
    2: "Cierre de comentarios (tarea 1)",
    20: "Revisión PostgreSQL",
    43: "API FIS organizaciones",
    44: "API FIS organizaciones",
    45: "API FIS datos/documentos",
}

COMENTARIOS = {
    19: "Q4: EF Core y migraciones en HapagPortal.DatabaseMigrations.",
    20: "Q4: EF Core y migraciones en HapagPortal.DatabaseMigrations.",
    28: "Permitir que el Portal avance mientras la integración definitiva con Nexus no esté disponible. "
        "Cubierto por adaptadores Dummy (6b) y simulador HTTP (6c).",
    37: "Resuelto por Q3: el portal consulta a Nexus vía API bajo demanda con caché corta; Nexus no actualiza "
        "tablas del portal.",
    42: "Q3 y Q6: contrato CT-FIS (docs/integraciones/contratos/fis.openapi.yaml).",
    43: "Q3 y Q6: contrato CT-FIS (docs/integraciones/contratos/fis.openapi.yaml).",
    44: "Q3 y Q6: contrato CT-FIS (docs/integraciones/contratos/fis.openapi.yaml).",
}

# Función NexusV2 (prefijo de la columna A) de las tareas 29–41.
TAREA_FUNCION = {
    29: "Creditos", 30: "Exenciones", 31: "Confirmacion Pago", 32: "Cambio de Almacen", 33: "Generar TATC",
    34: "Generar CLD", 35: "Conexion Portal", 36: "API Estado de Cuenta", 37: "Datos", 38: "Cuenta Admin",
    39: "TC", 40: "Tarifa", 41: "Counter",
}

# Resto de tareas originales: (R, apoyo) propuestos según Q5.
PROPUESTAS = [
    ((1, 2, 3, 4), ["Katu"], ["Kari"]),
    ((5, 6, 7, 8, 9, 10, 12), ["Andrés"], ["Jorge"]),
    ((11,), ["Jorge"], ["Andrés"]),
    ((13, 14, 16), ["Katu"], ["Andrés"]),
    ((15,), ["Jorge"], ["Andrés"]),
    ((17,), ["Andrés"], ["Jorge"]),
    ((18,), ["Jorge"], ["Katu"]),
    (tuple(range(19, 25)), ["Andrés"], ["Jorge"]),
    ((25,), ["Lucho"], ["Katu"]),
    ((26,), ["Lucho"], ["Jorge"]),
    ((27,), ["Jorge"], ["Katu"]),
    ((42, 43, 44, 46), ["Diego"], ["Jorge"]),
    ((45,), ["Diego"], ["Andrés"]),
    ((47,), ["Fer"], ["Ricardo"]),
    ((48, 50, 51, 53), ["Fer"], ["Jorge"]),
    ((49, 52, 55), ["Ricardo"], ["Jorge"]),
    ((54,), ["Lucho"], ["Fer", "Ricardo"]),
    ((56, 57, 61, 63, 64, 65, 66, 67, 68, 69, 70, 71, 72, 73, 74, 75), ["Kari"], ["Cami/Mati"]),
    ((58, 59, 60, 62, 76, 77, 78, 79, 80, 81), [LEGAL], ["Kari"]),
    ((82,), ["Andrés"], ["Kari"]),
    ((83, 85), ["Cami/Mati"], ["Fer"]),
    ((84,), ["Cami/Mati"], ["Andrés"]),
    ((86, 87, 88), ["Fer"], ["Ricardo"]),
    ((89,), ["Lucho"], ["Cami/Mati"]),
    ((90,), ["Andrés"], ["Lucho"]),
    ((91,), ["Cami/Mati"], ["Kari"]),
    ((92, 93), ["Andrés"], ["Cami/Mati"]),
    ((94, 95, 96, 98), ["Katu"], ["Cami/Mati"]),
    ((97, 99, 100), ["Andrés"], ["Jorge"]),
    ((101,), ["Andrés"], ["Katu"]),
    ((102, 103, 104, 105, 106), ["Diego"], ["RPX"]),
    ((107, 108, 109, 110, 111), [SEGURIDAD], ["Jorge"]),
    ((112, 113, 114, 115), [SEGURIDAD], ["Andrés"]),
    ((116, 117), ["Fer"], ["Jorge"]),
    ((118,), ["Fer"], ["Andrés"]),
]

FICHA_TAREA = {
    17: "NF-25", 18: "NF-25",
    25: "Cap. 15", 26: "Cap. 15", 27: "Cap. 15", 28: "Cap. 15; M4-01, M4-03, M5-05, M8-01",
    29: "M4-03, M5-07, M8-02", 30: "M4-01, M4-02", 31: "M5-06, NF-04", 32: "M3-04, M3-05, M3-06",
    33: "M2-09", 34: "M6-07, M3-16", 35: "Cap. 15", 36: "M7-03", 37: "Cap. 15", 38: "M8-06", 39: "M5-05",
    40: "M8-01", 41: "M8-09",
    42: "M8-04, M1-07", 43: "M2-06, M2-07, M2-09, M3-03", 44: "M2-02, M6-05", 45: "M2-06, M2-07, M2-09, M3-03",
    46: "M2-06, M2-07, M2-09, M3-03",
    47: "M5-01, M5-03", 48: "M5-01, M5-03", 49: "M5-01, M5-03", 50: "M7-01, M5-09", 51: "M7-01, M5-09",
    52: "M5-03", 53: "M5-03", 54: "M5-03, M5-06", 55: "M5-03",
    76: "M1-07, M6-06", 77: "M1-07",
    78: "M4-04, M6-06", 79: "M4-04, M6-06", 80: "M4-04, M6-06", 81: "M4-04, M6-06", 82: "M4-04, M6-06",
    83: "M3-01", 84: "M3-01", 85: "M3-01", 86: "M3-19", 87: "M3-19", 88: "M3-19",
    89: "M4-01, M4-02", 90: "M4-01, M4-02",
    91: "M8-09", 92: "M8-09", 93: "M8-09",
    94: "M1-25, M1-26", 95: "M1-25, M1-26", 96: "M1-25, M1-26", 97: "M1-25", 98: "M1-25", 99: "M1-25",
    100: "M1-25, M1-26", 101: "M1-25, M1-26",
    107: "NF-09", 108: "NF-09", 109: "NF-09", 110: "NF-09",
    111: "NF-05, NF-07", 112: "M1-02, M1-11", 113: "NF-07", 114: "NF-07", 115: "NF-07",
    116: "M5-03", 117: "M5-03", 118: "M5-03",
}
for _i in range(56, 76):
    FICHA_TAREA.setdefault(_i, "M1-07, M1-08")

CONTRATO_TAREA = {
    25: "CT-NEXUS", 26: "CT-NEXUS", 27: "CT-NEXUS", 28: "CT-NEXUS (fases 6b y 6c)", 29: "CT-NEXUS",
    30: "CT-NEXUS", 31: "CT-DEP", 32: "CT-MERC, CT-TATC", 33: "CT-TATC", 34: "CT-TATC", 35: "CT-NEXUS, CT-NAVE",
    36: "CT-NEXUS", 37: "CT-NEXUS", 39: "CT-NEXUS", 40: "CT-NEXUS",
    42: "CT-FIS", 43: "CT-FIS", 44: "CT-FIS", 45: "CT-FIS", 46: "CT-FIS",
    47: "CT-KHIPU, CT-BCH, CT-SANT, CT-BCI", 48: "CT-KHIPU, CT-BCH, CT-SANT, CT-BCI",
    49: "CT-KHIPU, CT-BCH, CT-SANT, CT-BCI", 50: "CT-DBNET", 51: "CT-DBNET", 52: "CT-BCH, CT-SANT",
    53: "CT-KHIPU", 54: "CT-DEP", 55: "CT-BCH, CT-SANT, CT-BCI",
    89: "CT-NEXUS", 90: "CT-NEXUS", 116: "CT-BCI", 117: "CT-BCI", 118: "CT-BCI",
}

# Tareas de nuevo desarrollo (lista cerrada). ID -> fases del plan cuyo inicio y fin dan las fechas.
# Solo estas filas llevan Fecha inicio y Fecha objetivo (config.CALENDARIO); su origen es "Plan Portal 2.0".
TAREAS_PLAN = {
    28: ("6b", "6c"),    # API Dummy de Nexus: adaptadores Dummy (6b) y simulador HTTP (6c)
    119: ("5a", "5a"),   # M11-01
    120: ("5b", "5b"),   # M11-02
    121: ("5b", "5b"),   # M11-03
    122: ("5c", "5c"),   # M11-04
    123: ("5c", "5c"),   # M11-05
    124: ("5c", "5c"),   # M11-06
    126: ("5a", "5a"),   # M11-08
    129: ("1", "1"),     # ficha M2-10
    130: ("1", "1"),     # ficha M8-09
}


def fechas_plan(tarea_id):
    inicio, fin = TAREAS_PLAN[tarea_id]
    return (datetime.date.fromisoformat(config.CALENDARIO[inicio][0]),
            datetime.date.fromisoformat(config.CALENDARIO[fin][1]))


def responsable_plan(tarea_id):
    fases = TAREAS_PLAN[tarea_id]
    return celda_responsable(EQUIPO[1], apoyo_fases(dict.fromkeys(fases))), EQUIPO[0], ORIGENES[1]


M11 = [
    ("M11-01", "Aplicar tokens de diseño a Bootstrap y eliminar colores sueltos (M11-01)", "UX", "Fase 5a del plan."),
    ("M11-02", "Selector de idioma ES/EN en caliente con Transloco (M11-02)", "UX", "Fase 5b del plan."),
    ("M11-03", "Formatos locales por país: fechas, números y montos (M11-03)", "UX", "Fase 5b del plan (Q8)."),
    ("M11-04", "Conformidad WCAG 2.2 AA en las pantallas principales (M11-04)", "UX",
     "Fase 5c del plan (Q9): ng lint, axe en ES y EN, Lighthouse y prueba manual."),
    ("M11-05", "Operación por teclado y foco visible (M11-05)", "UX", "Fase 5c del plan."),
    ("M11-06", "Anuncios dinámicos aria-live en pagos, cargas y errores (M11-06)", "UX", "Fase 5c del plan."),
    ("M11-07", "Tema claro/oscuro y reducción de movimiento (M11-07)", "UX",
     "Reducción de movimiento en 5c; selector de tema pendiente."),
    ("M11-08", "Navegadores y dispositivos soportados: browserslist y aviso de navegador no soportado (M11-08)", "UX",
     "Fase 5a del plan (NF-20, NF-21)."),
]


def tareas_nuevas(raci, fichas, cobertura):
    """Tareas nuevas en orden fijo desde el ID 119. Las de TAREAS_PLAN ocupan IDs estables."""
    nuevas = []

    def agregar(**t):
        t["id"] = 119 + len(nuevas)
        nuevas.append(t)

    modulo_m11 = "Interfaz, accesibilidad e idiomas (M11)"
    for fid, texto, tipo, coment in M11:
        agregar(clave=fid, modulo=modulo_m11, tarea=texto, dependencia="—", tipo=tipo, prioridad="Alta",
                comentario=coment, ficha=fid, contrato="—", estado="Pending", resp=(["Andrés"], ["Katu"]))
    agregar(clave="tema", modulo=modulo_m11, tarea="Selector de tema claro/oscuro en Angular (tokens oscuros ya declarados)",
            dependencia="M11-01", tipo="UX", prioridad="Media", comentario="Selector de tema pendiente (fuera de las fases 5a–5c).",
            ficha="M11-07", contrato="—", estado="Pending", resp=(["Andrés"], ["Katu"]))
    agregar(clave="localizacion", modulo=modulo_m11,
            tarea="Localizar ES/EN correos, PDF y asistente M10 (Fase 2 de Q8)", dependencia="M11-02", tipo="Desarrollo",
            prioridad="Media", comentario="Q8: en Fase 2 del portal; incluye ProblemDetails del backend.",
            ficha="M11-02, M10-05", contrato="—", estado="Pending", resp=(["Katu"], ["Andrés", "Kari"]))
    for fid, texto in (("M2-10", "Redactar la ficha M2-10 Consulta de plazos documentales por nave en la especificación v4"),
                       ("M8-09", "Redactar la ficha M8-09 Counter Bolivia/Ultramar en la especificación v4")):
        agregar(clave=fid, modulo="Especificación v4", tarea=texto, dependencia="Q7", tipo="Definición", prioridad="Media",
                comentario="Q7. Redactada en Portal_2.0_Especificacion_Funcional_v4.docx; pendiente de validación de Katu y Kari.",
                ficha=fid, contrato="CT-NEXUS" if fid == "M2-10" else "—", estado="In Progress")
    for fila, nombre, R, A, S, origen in raci:
        if origen != "Propuesta":
            continue
        clave = clave_funcion(nombre, NEXUS_COLUMNAS)
        _, ficha, _, contrato = NEXUS_COLUMNAS[clave]
        agregar(clave=f"raci-{clave}", modulo="Nexus",
                tarea=f"Confirmar el RACI propuesto de la función {clave} (NexusV2_v4)", dependencia="—", tipo="Definición",
                prioridad="Media", comentario="RACI propuesto en Funcionalidades NexusV2_v4.xlsx (S Jorge, C Andrés, I Katu/Kari).",
                ficha=ficha, contrato=contrato, estado="Pending", resp=(R.split(" / "), [A]))
    for cid, archivo, sistema in CONTRATOS:
        resp = responsable_contrato(archivo)
        agregar(clave=f"validar-{cid}", modulo="Contratos de integración",
                tarea=f"Validar contrato {cid} con {resp}", dependencia="Fase 6a", tipo="Integración", prioridad="Alta",
                comentario=f"Contrato en estado PROPUESTA (docs/integraciones/{archivo}). El plazo depende del responsable "
                           f"de Hapag-Lloyd.", ficha="Cap. 15", contrato=cid, estado="Pending",
                resp_contrato=resp)
    for sistema, cid, real in SISTEMAS_MODO:
        resp = dict((c[0], responsable_contrato(c[1])) for c in CONTRATOS)[cid]
        nota = ("Requiere contrato Validado y docs/integraciones/checklist-dummy-a-real.md."
                if real else
                "Sin cliente Real en el plan: requiere implementarlo, contrato Validado y el checklist Dummy→Real.")
        agregar(clave=f"conmutar-{sistema}", modulo="Integraciones Dummy→Real",
                tarea=f"Conmutar {sistema} Dummy→Real (Integrations:{sistema}:Mode)", dependencia=f"Validar contrato {cid}",
                tipo="Integración", prioridad="Media", comentario=nota, ficha="Cap. 15", contrato=cid, estado="Pending",
                resp=(["Jorge"], celda_desde_contrato(resp)[0].split(" (R)")[0]))
    for fid, titulo, fase in fichas:
        if fase_de(fase) != "1" or fid.startswith("M11") or cobertura.get(fid, ("N",))[0] != "N":
            continue
        agregar(clave=f"brecha-{fid}", modulo="Brechas de cobertura Fase 1",
                tarea=f"Implementar {fid} {titulo} (sin código en hapag-portal)", dependencia="—", tipo="Desarrollo",
                prioridad="Alta", comentario="Brecha N según Matriz_Fichas_Linea_Base_v4.xlsx.", ficha=fid,
                contrato=contrato_ficha(fid), estado="Pending", resp=(responsable_brecha(fid), EQUIPO[1]))
    esperado = {119 + i: f"M11-0{i + 1}" for i in range(8)}
    esperado.update({129: "M2-10", 130: "M8-09"})
    for t in nuevas:
        if t["id"] in esperado and t["clave"] != esperado[t["id"]]:
            raise SystemExit(f"ID {t['id']} no corresponde a {esperado[t['id']]} ({t['clave']})")
    return nuevas


RESP_BRECHA = {"M1": "Katu", "M2": "Cami/Mati", "M3": "Cami/Mati", "M4": "Lucho", "M5": "Fer", "M6": "Cami/Mati",
               "M7": "Fer", "M8": "Jorge", "M9": "Fer", "M10": "Katu", "NF": "Andrés"}


def responsable_brecha(fid):
    return [RESP_BRECHA[fid.split("-")[0]]]


def resolver_responsable(tarea_id, tarea, raci_por_funcion):
    """(celda, área, origen) según Q5."""
    if tarea_id in TAREAS_PLAN:
        return responsable_plan(tarea_id)
    if tarea_id in TAREA_FUNCION:
        _, _, R, A, S, origen = raci_por_funcion[TAREA_FUNCION[tarea_id]]
        r = R.split(" / ")
        apoyo = (S or A).split(" / ")
        return celda_responsable(r, apoyo), area_de(r[0]), ORIGENES[0] if origen == "RACI NexusV2" else ORIGENES[2]
    if tarea and "resp_contrato" in tarea:
        return celda_desde_contrato(tarea["resp_contrato"])
    if tarea and "resp" in tarea:
        r, apoyo = tarea["resp"]
        return celda_responsable(r, apoyo), area_de(r[0]), origen_propuesta(r)
    for ids, r, apoyo in PROPUESTAS:
        if tarea_id in ids:
            return celda_responsable(r, apoyo), area_de(r[0]), origen_propuesta(r)
    raise SystemExit(f"Tarea {tarea_id} sin regla de responsable")


def agregar_comentario(actual, texto):
    if not actual:
        return texto
    if texto in actual:
        return actual
    return actual.rstrip() + ("" if actual.rstrip().endswith(".") else ".") + " " + texto


def pendientes():
    raci = raci_v4()
    raci_por_funcion = {}
    for fila in raci:
        clave = clave_funcion(fila[1], TAREA_FUNCION.values())
        if clave:
            raci_por_funcion[clave] = fila
    fichas = fichas_v4()
    cobertura = matriz_base()
    nuevas = tareas_nuevas(raci, fichas, cobertura)

    asignadas = [i for ids, _, _ in PROPUESTAS for i in ids] + list(TAREA_FUNCION) + [28]
    if sorted(asignadas) != list(range(1, 119)):
        raise SystemExit("PROPUESTAS + TAREA_FUNCION + TAREAS_PLAN no cubren exactamente las tareas 1–118")

    wb = openpyxl.load_workbook(ORIG_PENDIENTES)
    ws = wb[HOJA_PENDIENTES]
    plantilla = {c: copy.copy(ws.cell(5, c)._style) for c in range(1, 12)}

    cabeceras = ["Área responsable", "Origen responsable", "Ficha spec v4", "Contrato de integración"]
    for i, cab in enumerate(cabeceras, start=12):
        c = ws.cell(4, i, cab)
        copiar_estilo(ws.cell(4, 11), c)
        for f in (1, 2):
            copiar_estilo(ws.cell(f, 11), ws.cell(f, i))

    # Filas originales (5..122): IDs 1..118.
    for r in range(5, ws.max_row + 1):
        tid = ws.cell(r, 1).value
        if not isinstance(tid, int):
            continue
        for rango, modulo in MODULOS_NUEVOS.items():
            if tid in rango:
                ws.cell(r, 2, modulo)
        if tid in TEXTOS:
            ws.cell(r, 3, TEXTOS[tid])
        if tid in DEPENDENCIAS:
            ws.cell(r, 5, DEPENDENCIAS[tid])
        if tid in COMENTARIOS:
            ws.cell(r, 11, COMENTARIOS[tid])
        if 102 <= tid <= 106:
            ws.cell(r, 11, agregar_comentario(ws.cell(r, 11).value, f"{FUERA_ALCANCE_MACROS}."))
        celda, area, origen = resolver_responsable(tid, None, raci_por_funcion)
        ws.cell(r, 4, celda)
        valores = [area, origen, FICHA_TAREA.get(tid, "—"), CONTRATO_TAREA.get(tid, "—")]
        for i, v in enumerate(valores, start=12):
            c = ws.cell(r, i, v)
            copiar_estilo(ws.cell(r, 11), c)

    # Tareas nuevas desde el ID 119.
    fila = ws.max_row + 1
    for t in nuevas:
        celda, area, origen = resolver_responsable(t["id"], t, raci_por_funcion)
        valores = [t["id"], t["modulo"], t["tarea"], celda, t["dependencia"], t["estado"], t["tipo"], t["prioridad"],
                   None, None, t["comentario"], area, origen, t["ficha"], t["contrato"]]
        for i, v in enumerate(valores, start=1):
            c = ws.cell(fila, i, v)
            c._style = copy.copy(plantilla[min(i, 11)])
        fila += 1
    ultima = fila - 1

    # Fechas: solo TAREAS_PLAN; el resto queda vacío con "Sin fecha en fuentes".
    for r in range(5, ultima + 1):
        tid = ws.cell(r, 1).value
        if tid in TAREAS_PLAN:
            inicio, fin = fechas_plan(tid)
            for col, v in ((9, inicio), (10, fin)):
                c = ws.cell(r, col, v)
                c.number_format = FORMATO_FECHA
            ws.cell(r, 4, responsable_plan(tid)[0])
            ws.cell(r, 12, EQUIPO[0])
            ws.cell(r, 13, ORIGENES[1])
        else:
            ws.cell(r, 9).value = None
            ws.cell(r, 10).value = None
            ws.cell(r, 11, agregar_comentario(ws.cell(r, 11).value, f"{SIN_FECHA}."))

    ws["A2"] = f"Consolidado v4 | Actualizado {datetime.date.today().strftime('%d-%m-%Y')}"
    rehacer_merge(ws, "A1:K1", "A1:O1")
    rehacer_merge(ws, "A2:K2", "A2:O2")
    for col, ancho in zip("LMNO", [22, 26, 22, 26]):
        ws.column_dimensions[col].width = ancho

    tabla = ws.tables["PortalTasksTable"]
    tabla.ref = f"A4:O{ultima}"
    tabla.autoFilter.ref = tabla.ref
    for i, cab in enumerate(cabeceras, start=12):
        tabla.tableColumns.append(TableColumn(id=i, name=cab))
    ampliar_cf(ws, ultima)
    ampliar_dv(ws, ultima)
    dv = DataValidation(type="list", formula1='"' + ",".join(ORIGENES) + '"', allow_blank=False)
    dv.add(f"M5:M{ultima}")
    ws.add_data_validation(dv)

    resumen(wb, ws, ultima)
    wb.calculation.fullCalcOnLoad = True
    salida = config.ruta_v4("Portal_2_0_Pendientes_v4.xlsx")
    wb.save(salida)
    return salida, ultima - 4, nuevas


def resumen(wb, ws_tareas, ultima):
    rs = wb["Resumen"]
    hoja = f"'{HOJA_PENDIENTES}'"
    cab_a, cab_b, val_a, val_b = (copy.copy(rs[c]._style) for c in ("A3", "B3", "A4", "B4"))
    cab_d, val_d, val_e = (copy.copy(rs[c]._style) for c in ("D3", "D4", "E4"))
    for fila in rs.iter_rows(min_row=3, max_row=rs.max_row):
        for c in fila:
            if c.row >= 11 or c.column >= 4:
                c.value = None

    # Totales por estado (A4:B9) quedan con sus fórmulas; el gráfico usa A5:B9.
    modulos = sorted({ws_tareas.cell(r, 2).value for r in range(5, ultima + 1)}, key=str.casefold)
    for i, titulo in enumerate(["Módulo", "Total", "Pendientes"]):
        c = rs.cell(3, 4 + i, titulo)
        c._style = copy.copy(cab_d)
    for j, m in enumerate(modulos, start=4):
        rs.cell(j, 4, m)._style = copy.copy(val_d)
        rs.cell(j, 5, f"=COUNTIF({hoja}!B:B,D{j})")._style = copy.copy(val_e)
        rs.cell(j, 6, f'=COUNTIFS({hoja}!B:B,D{j},{hoja}!F:F,"<>Completed")')._style = copy.copy(val_e)

    def bloque(fila, titulo, columna, valores):
        rs.cell(fila, 1, titulo)._style = copy.copy(cab_a)
        rs.cell(fila, 2, "Tareas")._style = copy.copy(cab_b)
        for k, v in enumerate(valores, start=fila + 1):
            rs.cell(k, 1, v)._style = copy.copy(val_a)
            rs.cell(k, 2, f"=COUNTIF({hoja}!{columna}:{columna},A{k})")._style = copy.copy(val_b)
        return fila + len(valores) + 2

    areas = sorted({ws_tareas.cell(r, 12).value for r in range(5, ultima + 1)}, key=str.casefold)
    siguiente = bloque(11, "Área responsable", "L", areas)
    bloque(siguiente, "Origen responsable", "M", ORIGENES)

    rs._charts = []
    grafico = BarChart()
    grafico.type = "col"
    grafico.grouping = "clustered"
    grafico.title = "Estado de tareas"
    grafico.legend = None
    grafico.add_data(Reference(rs, min_col=2, min_row=5, max_row=9), titles_from_data=False)
    # Categorías de texto como strRef (set_categories escribe numRef), igual que el gráfico original.
    grafico.series[0].cat = AxDataSource(strRef=StrRef(f"'Resumen'!$A$5:$A$9"))
    grafico.x_axis.delete = False
    grafico.y_axis.delete = False
    grafico.width = 15
    grafico.height = 8
    rs.add_chart(grafico, "H3")


# ================================================================ Comparación

HOJAS_ELIMINADAS = ["Alcance POC", "Texto fuente", "Evidencias"]
HOJAS_FASE = [("Primera fase", "1", False), ("Segunda fase", "2", False), ("Fase 0 - Revisión", "0", False),
              ("No funcionales", None, True)]
COLOR_COBERTURA = {"C": "FFC6EFCE", "P": "FFFFEB9C", "N": "FFFFC7CE"}

# Ficha con cobertura C o P -> (qué existe, brecha, archivo/pantalla, evidencia file:line en el commit de línea base).
DETALLE = {
    "M1-01": ("Shell con navbar, sidebar y footer; variables de color en styles.scss.",
              "Los overrides de Bootstrap no se aplican y no hay criterios medibles de consistencia (M11-01).",
              "Shell de la aplicación", "frontend/src/app/app.html:1-17; frontend/src/styles.scss:2-40"),
    "M1-02": ("CRUD de usuarios global con el permiso users.manage.",
              "Endpoints no acotados al cliente; la creación no asigna ClientId; la UI no llama a update ni a set-active.",
              "Administración > Usuarios",
              "backend/src/HapagPortal.WebApi/Controllers/V1/UsersController.cs:17-59; "
              "frontend/src/app/core/services/admin-user.service.ts:19-35"),
    "M1-04": ("El país queda fijado por el cliente al registrarse y la navbar lo muestra en solo lectura.",
              "No hay selector de país.", "Navbar",
              "backend/src/HapagPortal.Domain/Entities/Client.cs:10; frontend/src/app/shared/components/navbar/navbar.html:25-30"),
    "M1-05": ("Dashboard con 3 contadores; el panel de país es texto fijo.",
              "Faltan indicadores por servicio y contenido dinámico por país.", "Dashboard",
              "frontend/src/app/features/dashboard/dashboard.ts:30-47; frontend/src/app/features/dashboard/dashboard.html:105-192"),
    "M1-07": ("Registro de organización para los tipos Client y CustomsAgent.",
              "La organización queda activa sin aprobación; no admite otros tipos ni adjuntar documentos.", "Registro",
              "backend/src/HapagPortal.Application/Auth/Register/RegisterCommandHandler.cs:20-117"),
    "M1-10": ("Recuperación de contraseña completa: respuesta uniforme, token de 1 h de un solo uso.",
              "El cierre de sesión es solo local; no hay endpoint de logout.", "Login / Recuperar contraseña",
              "backend/src/HapagPortal.Application/Auth/ForgotPassword/ForgotPasswordCommandHandler.cs:14-43; "
              "backend/src/HapagPortal.Application/Auth/ResetPassword/ResetPasswordCommandHandler.cs:16-55; "
              "frontend/src/app/core/services/auth.service.ts:81-87"),
    "M1-11": ("RBAC en el servidor con 8 roles y 14 permisos sembrados, aplicados con [HasPermission].",
              "Sin matriz por tipo de información ni por acción sobre el BL; roles.manage sin uso.", "—",
              "backend/src/HapagPortal.Infrastructure/Persistence/ApplicationDbContext.cs:108-165; "
              "backend/src/HapagPortal.Infrastructure/Authentication/PermissionResolver.cs:10-31; "
              "backend/src/HapagPortal.Infrastructure/Authentication/HasPermissionAttribute.cs:5-52"),
    "M1-23": ("Consulta de auditoría de solo lectura.", "Ningún código escribe AuditLogs en ejecución.",
              "Administración > Auditoría",
              "backend/src/HapagPortal.WebApi/Controllers/V1/AuditController.cs:15-23; "
              "backend/src/HapagPortal.Application/Audit/Search/SearchAuditQuery.cs:19-65"),
    "M1-25": ("Entidad Notification, NotificationsController y campana en la navbar.",
              "Solo emiten notificaciones los plazos y las transmisiones de Aduana.", "Navbar (campana)",
              "backend/src/HapagPortal.WebApi/Controllers/V1/NotificationsController.cs:15-43; "
              "backend/src/HapagPortal.Infrastructure/Notifications/NotificationPublisher.cs:17-51; "
              "frontend/src/app/shared/components/navbar/navbar.html:33-40"),
    "M2-06": ("Listado y detalle de los BL del propio cliente.",
              "Solo búsqueda por número exacto; sin booking ni filtros por columna.", "Embarques > Listado y detalle",
              "frontend/src/app/features/bill-of-lading/bl-list/bl-list.ts:34-73; "
              "frontend/src/app/features/bill-of-lading/bl-detail/bl-detail.ts:34-60; "
              "backend/src/HapagPortal.Application/BillsOfLading/Read/GetMyBLs/GetMyBLsQueryHandler.cs:39"),
    "M2-07": ("ShipmentType se muestra como columna Tipo.", "Sin filtro ni persistencia de la separación.",
              "Embarques > Listado", "frontend/src/app/features/bill-of-lading/bl-list/bl-list.html:56"),
    "M2-09": ("Consulta de BL por número, restringida al cliente propietario.", "No hay TATC.", "Embarques > Detalle",
              "backend/src/HapagPortal.Application/BillsOfLading/Read/GetByNumber/GetBLByNumberQueryHandler.cs:26"),
    "M3-01": ("Lectura genérica de los cargos almacenados por BL.",
              "Sin lectura Gate Out desde el origen ni aplicación de reglas de exención.", "Cargos locales",
              "backend/src/HapagPortal.Application/LocalCharges/Read/GetByBL/GetLocalChargesByBLQueryHandler.cs:19-43"),
    "M3-04": ("Solicitud individual de cambio de almacén con monto.",
              "Siempre cobra (sin rama gratuita); el formulario no envía motivo, contenedor ni teléfono.",
              "Cambio de almacén",
              "backend/src/HapagPortal.Application/WarehouseChanges/Commands/Create/CreateWarehouseChangeCommandHandler.cs:21-55; "
              "frontend/src/app/features/warehouse/warehouse.ts:153-159"),
    "M3-06": ("Endpoints GET warehouse-changes/my y GET warehouse-changes/{id}.",
              "Sin historial de estados ni RUT del pagador.", "Cambio de almacén > Mis solicitudes",
              "backend/src/HapagPortal.WebApi/Controllers/V1/WarehouseChangesController.cs:28-40"),
    "M3-12": ("Rectificación de manifiesto de Aduana.", "El cliente no puede solicitar ni pagar la corrección.", "Aduana",
              "backend/src/HapagPortal.Application/Customs/Amend/SubmitManifestAmendmentCommand.cs:16-21"),
    "M3-13": ("La transmisión del BL hijo exige el padre aceptado; existe la regla BL_HOUSE_IN.", "Sin cobro.", "Aduana",
              "backend/src/HapagPortal.Application/Customs/Transmit/TransmitBLCommand.cs:50-59"),
    "M3-14": ("DeadlineCalculator produce el estado Overdue.", "Sin cobro por matriz fuera de plazo.", "Aduana > Plazos",
              "backend/src/HapagPortal.Domain/Validation/DeadlineCalculator.cs:9-36"),
    "M4-01": ("Entidad local DemurrageExemption con CRUD de administración.",
              "Ningún flujo la consume y no se lee desde Nexus.", "Administración > Exenciones",
              "backend/src/HapagPortal.Domain/Entities/DemurrageExemption.cs:5-11"),
    "M4-03": ("Entidad local CreditClient con CRUD de administración.",
              "Ningún flujo la consume; no existe el recargo IPO.", "Administración > Créditos",
              "backend/src/HapagPortal.Domain/Entities/CreditClient.cs:5-15"),
    "M5-01": ("CreatePaymentCommand acepta varios Details.",
              "El handler no usa ChargeIds; no hay entidad ni vista de carro.", "Pagos > Nuevo pago",
              "backend/src/HapagPortal.Application/Payments/Create/CreatePaymentCommand.cs:6-14"),
    "M5-02": ("Cancelación de pagos no confirmados.", "Sin control de anulación de boletas por parte del cliente.", "Pagos",
              "backend/src/HapagPortal.Application/Payments/Commands/Cancel/CancelPaymentCommandHandler.cs:35-38"),
    "M5-03": ("Constantes Khipu, BankButton, Deposit y BCI; endpoint de medios de pago por país.",
              "El frontend usa una lista fija y no consume el endpoint.", "Pagos > Nuevo pago",
              "backend/src/HapagPortal.Domain/Constants/PaymentMethods.cs:12-15; "
              "backend/src/HapagPortal.Application/Config/Read/GetPaymentMethods/GetPaymentMethodsQueryHandler.cs:15-31; "
              "frontend/src/app/features/payments/payment-form/payment-form.ts:44-58"),
    "M5-04": ("Entidad Currency y endpoint de monedas.", "El pago usa una sola moneda, la del BL.", "Pagos",
              "backend/src/HapagPortal.Domain/Entities/Currency.cs:5-11; "
              "backend/src/HapagPortal.Application/Payments/Create/CreatePaymentCommandHandler.cs:82-84"),
    "M5-06": ("Campo Payment.DepositProofUrl.", "Sin endpoint ni UI de carga; no hay almacenamiento de archivos.", "—",
              "backend/src/HapagPortal.Domain/Entities/Payment.cs:22"),
    "M5-07": ("Método de pago CreditLine.", "No se valida contra CreditClient ni contra la condición de crédito de Nexus.",
              "Pagos > Nuevo pago",
              "backend/src/HapagPortal.Application/Payments/Create/CreatePaymentCommandHandler.cs:43; "
              "backend/src/HapagPortal.Application/Admin/CreditClients/Read/GetAll/GetAllCreditClientsQueryHandler.cs:22-31"),
    "M7-02": ("GET payments/my y GET receipts/my con sus pantallas.", "El PDF es placeholder y no hay RUT del pagador.",
              "Pagos > Historial; Recibos",
              "backend/src/HapagPortal.WebApi/Controllers/V1/PaymentsController.cs:43; "
              "backend/src/HapagPortal.WebApi/Controllers/V1/ReceiptsController.cs:54; "
              "backend/src/HapagPortal.Application/Receipts/Read/GetPdf/GetReceiptPdfQueryHandler.cs:31"),
    "M6-08": ("Tipos de orden de servicio RELEASE y DECONSOLIDATION.", "Sin emisión de la carta ni flujo de Bolivia.",
              "Órdenes de servicio", "frontend/src/app/features/service-orders/service-orders.ts:38-46"),
    "M8-05": ("Rutas admin/*: créditos, exenciones, usuarios, importación de BL, Aduana, plazos, auditoría y reportes.",
              "Falta el área completa descrita en la ficha.", "Administración", "frontend/src/app/app.routes.ts:118-164"),
    "M8-06": ("Administrador, SuperAdmin y Admin reciben todos los permisos.",
              "El rol expuesto es ADMIN/USER; no hay visibilidad total diferenciada.", "—",
              "backend/src/HapagPortal.Infrastructure/Authentication/PermissionResolver.cs:10-20; "
              "backend/src/HapagPortal.Application/Auth/Login/LoginCommandHandler.cs:55"),
    "M10-02": ("FAQ por país sembrada en BD, con búsqueda por texto.", "No hay asistente conversacional.", "FAQ",
               "backend/src/HapagPortal.WebApi/Controllers/V1/FAQsController.cs:13-36; frontend/src/app/features/faq/faq.ts:34-63"),
    "NF-01": ("Webhooks y recibo idempotentes: un estado terminal no tiene efectos.",
              "La creación de pago no tiene clave de idempotencia.", "—",
              "backend/src/HapagPortal.Application/Payments/Commands/Webhooks/KhipuWebhookCommandHandler.cs:40-41"),
    "NF-02": ("Constantes de estado de pago.", "Sin historial de transiciones; PendingVerification sin uso.", "—",
              "backend/src/HapagPortal.Domain/Constants/PaymentStatus.cs:5-10"),
    "NF-04": ("PaymentNumber único y ExternalReference.", "Sin TransactionId persistido.", "—",
              "backend/src/HapagPortal.Domain/Entities/Payment.cs:7-20"),
    "NF-05": ("Filtro por ClientId en cada handler.", "Sin filtro global.", "—",
              "backend/src/HapagPortal.Application/BillsOfLading/Read/GetByNumber/GetBLByNumberQueryHandler.cs:26"),
    "NF-07": ("AES-GCM para secretos y BCrypt para contraseñas.",
              "Sin HTTPS/HSTS en la línea base (se agregan en la fase 6c).", "—",
              "backend/src/HapagPortal.Infrastructure/Secrets/AesSecretProtector.cs:30-62; "
              "backend/src/HapagPortal.Infrastructure/Authentication/PasswordHasher.cs:10-13"),
    "NF-08": ("Payment guarda solo la referencia y el resultado; no almacena datos del instrumento.", "—", "—",
              "backend/src/HapagPortal.Domain/Entities/Payment.cs:5-22"),
    "NF-09": ("SecretCredential cifrado y SecretResolver.", "En la línea base solo hay tipos para SII y Aduana.", "—",
              "backend/src/HapagPortal.Domain/Entities/SecretCredential.cs:9-15; "
              "backend/src/HapagPortal.Infrastructure/Secrets/SecretResolver.cs:9-31"),
    "NF-12": ("Si la pasarela falla, el pago queda Failed.", "Sin timeout ni reintento en la línea base.", "—",
              "backend/src/HapagPortal.Application/Payments/Create/CreatePaymentCommandHandler.cs:183-190"),
    "NF-14": ("Columnas CreatedBy/ModifiedBy completadas por interceptor.", "Sin registro de operaciones por cliente.", "—",
              "backend/src/HapagPortal.Infrastructure/Persistence/Interceptors/AuditableEntityInterceptor.cs:27-55"),
    "NF-15": ("AuditLog con OldValues y NewValues.", "Ningún código escribe el valor anterior.", "—",
              "backend/src/HapagPortal.Domain/Entities/AuditLog.cs:10-11"),
    "NF-16": ("Tabla de auditoría y borrado lógico.", "Sin política de conservación.", "—",
              "backend/src/HapagPortal.Infrastructure/Persistence/Interceptors/AuditableEntityInterceptor.cs:48-51"),
    "NF-19": ("Importación de BL síncrona con resultado por fila.", "Sin procesamiento en segundo plano.",
              "Administración > Importar BL",
              "backend/src/HapagPortal.Application/BillsOfLading/Import/ImportBillsOfLadingCommandHandler.cs:25-119"),
    "NF-21": ("Grilla Bootstrap y table-responsive en las tablas.", "Un único breakpoint.", "—",
              "frontend/src/styles.scss:434"),
    "NF-22": ("UTC en auditoría.", "Sin huso horario de referencia por país.", "—",
              "backend/src/HapagPortal.Infrastructure/Persistence/Interceptors/AuditableEntityInterceptor.cs:31"),
    "NF-23": ("Interfaz solo en español.", "Sin LOCALE_ID ni i18n (M11-02, M11-03).", "—", "frontend/src/index.html:2"),
    "NF-24": ("Staging en Railway con datos demo.", "Sin ambiente de pruebas de integraciones.", "—",
              "RAILWAY-DEPLOY.md:30-57"),
    "NF-25": ("Dockerfile y migraciones automáticas al arrancar.", "Sin rollback.", "—",
              "backend/Dockerfile:1-16; backend/src/HapagPortal.WebApi/Program.cs:134"),
    "NF-26": ("Solo el endpoint /health.", "Sin monitoreo ni alertas.", "—", "backend/src/HapagPortal.WebApi/Program.cs:165"),
    "NF-27": ("ExceptionHandlingMiddleware con ILogger.",
              "Sin correlación ni registro estructurado de integraciones en la línea base (se agregan en la fase 6c).", "—",
              "backend/src/HapagPortal.WebApi/Middleware/ExceptionHandlingMiddleware.cs:25-37"),
}

ALERTAS = {
    "M1-17": "Fase 1 con confirmación de inclusión pendiente en las dependencias de la ficha.",
    "M3-02": "Resuelto por Q1: rige el encabezado (Fase 1); se quitó 'en su segunda fase'.",
    "M6-02": "Resuelto por Q1: Fase 2; 'primera etapa' pasa a 'primera entrega dentro de Fase 2, sin pago ni carro'.",
    "M3-17": "Resuelto por Q1: Fase 2; el Web Service lo desarrolla y administra Hapag-Lloyd.",
    "M5-07": "Resuelto por Q1: depende de M4-03/M8-02 (lectura de condición de crédito), no de M7-03.",
    "M6-07": "Resuelto por Q1: depende de M4-03/M8-02 (lectura de condición de crédito), no de M7-03.",
    "NF-27": "Resuelto por Q2: se eliminaron las referencias a fichas inexistentes; NF-27 queda autocontenida.",
}

FASE_PLAN_M11 = {"M11-01": "5a", "M11-08": "5a", "M11-02": "5b", "M11-03": "5b", "M11-04": "5c", "M11-05": "5c",
                 "M11-06": "5c", "M11-07": "5c (reducción de movimiento); selector de tema pendiente"}


def nombre_archivo(evidencia):
    return [os.path.basename(p.split(":")[0]) for p in evidencia.split("; ")]


def fila_comparacion(fid, titulo, fase, cobertura):
    nueva = fid not in cobertura
    cob = "N" if nueva else cobertura[fid][0]
    puertos = codigo_puertos(fid)
    if cob in ("C", "P"):
        if fid not in DETALLE:
            raise SystemExit(f"Ficha {fid} con cobertura {cob} sin detalle en DETALLE")
        existe, brecha, pantalla, evidencia = DETALLE[fid]
        codigo = ", ".join(nombre_archivo(evidencia))
        if puertos:
            codigo += "; " + puertos
    else:
        evidencia, pantalla = "—", "—"
        codigo = puertos or "—"
        if nueva:
            existe = "Ficha nueva de la especificación v4; sin código en la línea base."
            if fid in FASE_PLAN_M11:
                brecha = f"Se aborda en la fase {FASE_PLAN_M11[fid]} del plan Portal 2.0."
            else:
                brecha = "Implementar la ficha completa (Fase 2 del portal)."
        else:
            existe = "No existe código para esta ficha en hapag-portal."
            brecha = "Implementar la ficha completa."
            if puertos:
                existe += " Hay puerto de integración con adaptador Dummy, sin consumidor de negocio."
    if fid.startswith("M11"):
        alerta = "Ficha nueva v4 (DC1)."
    elif nueva:
        alerta = "Ficha nueva v4 (Q7)."
    else:
        alerta = ALERTAS.get(fid)
    referencia = f"Ficha nueva v4; {fase}" if nueva else f"Ficha {fid}; {fase}"
    return [fid, fid.split("-")[0], titulo, cob, codigo, existe, brecha, pantalla, evidencia, fase, alerta, referencia]


def comparacion():
    fichas = fichas_v4()
    cobertura = matriz_base()
    base, fecha = commit_base()
    wb = openpyxl.load_workbook(ORIG_COMPARACION)
    for nombre in HOJAS_ELIMINADAS:
        del wb[nombre]

    cabeceras = {4: "Cobertura hapag-portal (C/P/N)", 5: "Código relacionado", 8: "Archivo/pantalla en hapag-portal",
                 9: "Evidencia (file:line)"}
    nota = (f"Cobertura hapag-portal: C construido, P parcial, N no existe, medida en el commit de línea base {base} "
            f"({fecha}); las evidencias file:line son de ese commit. 'Código relacionado' agrega los puertos y "
            "adaptadores de integración de las fases 6b y 6c.")
    conteo = {}
    for hoja, fase, es_nf in HOJAS_FASE:
        ws = wb[hoja]
        estilos = [{c: copy.copy(ws.cell(r, c)._style) for c in range(1, 13)} for r in (5, 6)]
        alturas = [ws.row_dimensions[5].height, ws.row_dimensions[6].height]
        ws.delete_rows(5, ws.max_row - 4)
        ws["A1"] = f"{hoja} | Requerimientos vs hapag-portal"
        ws["A2"] = nota
        for col, texto in cabeceras.items():
            ws.cell(4, col, texto)
        filas = [f for f in fichas if (f[0].startswith("NF") if es_nf else (not f[0].startswith("NF") and fase_de(f[2]) == fase))]
        for i, (fid, titulo, fase_txt) in enumerate(filas):
            r = 5 + i
            for c, v in enumerate(fila_comparacion(fid, titulo, fase_txt, cobertura), start=1):
                celda = ws.cell(r, c, v)
                celda._style = copy.copy(estilos[i % 2][c])
            d = ws.cell(r, 4)
            d.fill = PatternFill("solid", fgColor=COLOR_COBERTURA[d.value])
            ws.row_dimensions[r].height = alturas[i % 2]
        ultima = 4 + len(filas)
        ws.auto_filter.ref = f"A4:L{ultima}"
        ws.column_dimensions["E"].width = 45
        ws.column_dimensions["I"].width = 60
        conteo[hoja] = (len(filas), ultima)

    resumen_comparacion(wb, conteo, fichas, cobertura, base, fecha)
    catalogo(wb, cobertura, fichas, base)
    wb.calculation.fullCalcOnLoad = True
    salida = config.ruta_v4("Comparacion_Requerimientos_Por_Fase_v4.xlsx")
    wb.save(salida)
    return salida, conteo


def resumen_comparacion(wb, conteo, fichas, cobertura, base, fecha):
    rs = wb["Resumen"]
    rs["A1"] = "Comparación hapag-portal vs requerimiento funcional v4"
    rs["A2"] = (f"Fuente: Portal_2.0_Especificacion_Funcional_v4.docx ({len(fichas)} fichas). Código: repositorio "
                f"hapag-portal, commit de línea base {base} ({fecha}). Cobertura según Matriz_Fichas_Linea_Base_v4.xlsx.")
    for col, texto in zip("ABCDEFGH", ["Grupo documental", "Total fichas", "C (construido)", "P (parcial)",
                                       "N (no existe)", "Fichas nuevas v4", "Relacionado con código existente",
                                       "% con código (C+P)"]):
        rs[f"{col}4"] = texto
    for fila, (hoja, _, _) in zip(range(5, 9), HOJAS_FASE):
        _, ultima = conteo[hoja]
        h = f"'{hoja}'"
        rs[f"A{fila}"] = hoja
        rs[f"B{fila}"] = f"=COUNTA({h}!A5:A{ultima})"
        rs[f"C{fila}"] = f'=COUNTIF({h}!D5:D{ultima},"C")'
        rs[f"D{fila}"] = f'=COUNTIF({h}!D5:D{ultima},"P")'
        rs[f"E{fila}"] = f'=COUNTIF({h}!D5:D{ultima},"N")'
        rs[f"F{fila}"] = f'=COUNTIF({h}!L5:L{ultima},"Ficha nueva v4*")'
        rs[f"G{fila}"] = f'=COUNTA({h}!A5:A{ultima})-COUNTIF({h}!E5:E{ultima},"—")'
        rs[f"H{fila}"] = f"=IFERROR((C{fila}+D{fila})/B{fila},0)"
    rs["A10"] = "Cómo interpretar Cobertura hapag-portal"
    leyenda = [
        ("C", "Construido: todos los criterios de la ficha tienen código en hapag-portal."),
        ("P", "Parcial: existe código relacionado y faltan criterios de la ficha."),
        ("N", "No existe: no hay código para la ficha. Incluye las fichas nuevas de la v4."),
        ("Nueva v4", "Ficha agregada en la especificación v4 (M11-01 a M11-08, M2-10 y M8-09); sin código en la línea base."),
    ]
    for fila, (a, b) in zip(range(11, 15), leyenda):
        rs[f"A{fila}"] = a
        rs[f"B{fila}"] = b

    rf = [f for f in fichas if not f[0].startswith("NF")]
    por_fase = {k: sum(1 for f in rf if fase_de(f[2]) == k) for k in ("1", "2", "0")}
    n_nf = len(fichas) - len(rf)
    n_f1 = [f[0] for f in fichas if fase_de(f[2]) == "1" and cobertura.get(f[0], ("N",))[0] == "N"
            and not f[0].startswith("M11")]
    p_f1 = [f[0] for f in fichas if fase_de(f[2]) == "1" and cobertura.get(f[0], ("N",))[0] == "P"]
    notas = [
        "La comparación mide el código del repositorio hapag-portal (Angular 21, .NET 9 y PostgreSQL) contra las fichas "
        "de la especificación v4. No usa ningún sitio externo como referencia.",
        f"Las {len(rf)} fichas funcionales se separan en {por_fase['1']} de Fase 1, {por_fase['2']} de Fase 2 y "
        f"{por_fase['0']} de Fase 0. Los {n_nf} no funcionales se muestran aparte, todos con encabezado Fase 1. "
        "El catálogo incluye 52 servicios y no se suma a las fichas.",
        "Las fichas nuevas de la v4 (M11-01 a M11-08, M2-10 y M8-09) se marcan N: no tenían código en la línea base.",
        "Los puertos, adaptadores Dummy y clientes Real de las fases 6b y 6c figuran en 'Código relacionado'. No cambian "
        "la cobertura porque ningún flujo de negocio los consume todavía. Todos los contratos de integración están en "
        "estado PROPUESTA.",
        "Alertas de fase resueltas por Q1: M3-02 = Fase 1; M6-02 = Fase 2 (primera entrega sin pago ni carro); M3-17 = "
        "Fase 2. M5-07 y M6-07 dependen de M4-03/M8-02 (lectura de condición de crédito), no de M7-03.",
        "Capítulo 15 (integración): APIs de Hapag-Lloyd sin acceso directo a BD; FIS y Data Lake por API (Q3); entrada "
        "transitoria de FIS por POST bills-of-lading/import (Q6). Inventario y contratos en docs/integraciones/.",
        f"Brechas de Fase 1 (RF y NF): {len(n_f1)} fichas sin código (N, sin contar M11) y {len(p_f1)} parciales (P). Cada ficha N "
        "tiene su tarea en el módulo 'Brechas de cobertura Fase 1' de Pendientes_v4.",
    ]
    for fila, texto in zip(range(16, 23), notas):
        rs[f"A{fila}"] = texto


def catalogo(wb, cobertura, fichas, base):
    ws = wb["Catálogo servicios"]
    ids = {f[0] for f in fichas}
    ws["A2"] = ("52 servicios del capítulo 3. Complementan las fichas de la especificación v4; no se suman como "
                "requisitos adicionales. La cobertura corresponde a la ficha relacionada en hapag-portal, no a "
                "certificación por país.")
    ws["F4"] = "Cobertura hapag-portal de la ficha"
    ws["H4"] = "Límite de validación"
    for r in range(5, ws.max_row + 1):
        if not ws.cell(r, 1).value:
            continue
        relacionadas = [x.strip() for x in str(ws.cell(r, 5).value or "").split(",") if x.strip()]
        if relacionadas:
            partes = []
            for fid in relacionadas:
                if fid in ids:
                    partes.append(f"{fid}: {cobertura.get(fid, ('N',))[0]}")
                else:
                    partes.append(f"{fid}: sin ficha v4")
            ws.cell(r, 6, "; ".join(partes))
        else:
            ws.cell(r, 6, "Sin ficha relacionada")
        ws.cell(r, 8, f"Cobertura del código de hapag-portal (commit {base}) para la ficha relacionada. No se "
                      "ejecutaron pagos ni integraciones por país.")


# ================================================================ main

def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    for p in (ORIG_PENDIENTES, ORIG_NEXUS, ORIG_COMPARACION, ESPEC_V4, MATRIZ_BASE):
        if not os.path.exists(p):
            raise SystemExit(f"No existe {p}")
    s_nexus = nexus()
    s_pend, total, nuevas = pendientes()
    s_comp, conteo = comparacion()
    origenes = {}
    wb = openpyxl.load_workbook(s_pend, read_only=True)
    for fila in wb[HOJA_PENDIENTES].iter_rows(min_row=5, values_only=True):
        if fila[0] is not None:
            origenes[fila[12]] = origenes.get(fila[12], 0) + 1
    print(f"NexusV2: {s_nexus}")
    print(f"Pendientes: {total} tareas (118 originales + {len(nuevas)} nuevas, IDs 119–{118 + len(nuevas)}) -> {s_pend}")
    print("  por origen: " + ", ".join(f"{k} {v}" for k, v in origenes.items()))
    print(f"  con fecha (TAREAS_PLAN): {sorted(TAREAS_PLAN)}")
    print(f"Comparación: {', '.join(f'{h} {n}' for h, (n, _) in conteo.items())} -> {s_comp}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
