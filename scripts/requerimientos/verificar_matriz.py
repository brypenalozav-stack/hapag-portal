"""Fase 3: verifica la Matriz de trazabilidad (Markdown y Word). Sale con 1 si algo falla.

- 134 filas, con los IDs de las fichas de la especificación v4 en el mismo orden, en el .md y en el .docx.
- Cada ficha de Fase 1 tiene al menos una tarea de Pendientes_v4.
- Los contratos citados existen en la tabla de la Fase 6a.
- En las filas con contrato, cada línea de "Estado contrato" empieza por "PROPUESTA" o por "Validado" y cada
  contrato de la fila tiene su estado.
- En las filas con un contrato que tiene puerto, "Código" cita un Dummy*.cs; si el contrato es CT-NEXUS, CT-FIS,
  CT-KHIPU, CT-BCH, CT-DBNET o CT-TRACK, cita además un Http*.cs. Todo archivo citado existe en backend/src.
- El .docx es horizontal y su tabla repite la fila de cabecera.
- El .md y el .docx tienen el mismo número de filas.
- 0 términos prohibidos (terminos.py) en ambos archivos.
"""
import os
import re
import sys

import docx
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402
import docx_utils  # noqa: E402
import terminos  # noqa: E402
from linea_base import DOCS, REPO  # noqa: E402

MD = os.path.join(DOCS, "matriz-trazabilidad.md")
DOCX = os.path.join(config.V4, "Matriz_Trazabilidad_v4.docx")
ESPEC_V4 = os.path.join(config.V4, "Portal_2.0_Especificacion_Funcional_v4.docx")
BACKEND_SRC = os.path.join(REPO, "backend", "src")

# Tabla de contratos de la Fase 6a del plan.
CONTRATOS_6A = {"CT-NEXUS", "CT-FIS", "CT-KHIPU", "CT-BCH", "CT-SANT", "CT-BCI", "CT-DBNET", "CT-TRACK", "CT-SIGN",
                "CT-STORAGE", "CT-MERC", "CT-TATC", "CT-NAVE", "CT-DEP", "CT-DISP"}
CON_PUERTO = {"CT-NEXUS", "CT-FIS", "CT-KHIPU", "CT-BCH", "CT-SANT", "CT-BCI", "CT-DBNET", "CT-TRACK", "CT-SIGN",
              "CT-STORAGE"}
CON_HTTP = {"CT-NEXUS", "CT-FIS", "CT-KHIPU", "CT-BCH", "CT-DBNET", "CT-TRACK"}

COLUMNAS = ["Ficha", "Fase", "Tareas Pendientes_v4", "Función NexusV2", "Contrato", "Estado contrato",
            "Código hapag-portal", "Cobertura"]
RE_CONTRATO = re.compile(r"\bCT-[A-Z]+\b")
RE_ARCHIVO = re.compile(r"[\w./-]+\.cs\b")
SIN_TAREA = "Sin tarea"


def lineas_celda(celda):
    return [x.strip() for x in re.split(r"<br>|\n", celda) if x.strip()]


def tabla_md():
    with open(MD, encoding="utf-8") as f:
        texto = f.read()
    filas, dentro = [], False
    for l in texto.splitlines():
        if l.startswith("| Ficha |"):
            dentro = True
            filas.append([c.strip() for c in re.split(r"(?<!\\)\|", l.strip())[1:-1]])
            continue
        if dentro:
            if not l.startswith("|"):
                break
            celdas = [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", l.strip())[1:-1]]
            if all(re.fullmatch(r"-+", c) for c in celdas):
                continue
            filas.append(celdas)
    return texto, filas


def tabla_docx():
    d = docx.Document(DOCX)
    tabla = next((t for t in d.tables if t.rows[0].cells[0].text.strip() == "Ficha"), None)
    filas = []
    repite = False
    if tabla is not None:
        filas = [["\n".join(p.text for p in c.paragraphs).strip() for c in fila.cells] for fila in tabla.rows]
        trpr = tabla.rows[0]._tr.trPr
        repite = trpr is not None and trpr.find(qn("w:tblHeader")) is not None
    horizontal = d.sections[0].orientation == WD_ORIENT.LANDSCAPE and \
        d.sections[0].page_width > d.sections[0].page_height
    texto = "\n".join("".join(x.text or "" for x in p.iter(qn("w:t"))) for p in d.element.body.iter(qn("w:p")))
    texto += "\n" + "\n".join(p.text for s in d.sections for p in s.footer.paragraphs)
    return texto, filas, repite, horizontal


def verificar_filas(nombre, filas, ids_spec, errores):
    if not filas:
        errores.append(f"{nombre}: no hay tabla con cabecera 'Ficha'")
        return {}
    if filas[0] != COLUMNAS:
        errores.append(f"{nombre}: columnas {filas[0]} (se esperaban {COLUMNAS})")
        return {}
    cuerpo = filas[1:]
    if len(cuerpo) != 134:
        errores.append(f"{nombre}: {len(cuerpo)} filas (se esperaban 134)")
    ids = [f[0].split(" ")[0] for f in cuerpo]
    if ids != ids_spec:
        errores.append(f"{nombre}: los IDs de ficha no coinciden con la especificación v4 en orden")

    sin_tarea_f1, citados = [], set()
    n_ct = n_codigo = 0
    for f in cuerpo:
        ficha, fase, tareas, _, contrato, estado, codigo, cobertura = f
        fid = ficha.split(" ")[0]
        if fase not in ("0", "1", "2"):
            errores.append(f"{nombre}: {fid} con fase {fase!r}")
        if cobertura not in ("C", "P", "N"):
            errores.append(f"{nombre}: {fid} con cobertura {cobertura!r}")
        if fase == "1" and (not tareas.strip() or tareas.strip() == SIN_TAREA):
            sin_tarea_f1.append(fid)
        cts = RE_CONTRATO.findall(contrato)
        for c in cts:
            citados.add(c)
            if c not in CONTRATOS_6A:
                errores.append(f"{nombre}: {fid} cita el contrato {c}, que no está en la tabla de la Fase 6a")
        if not cts:
            continue
        n_ct += 1
        lineas = lineas_celda(estado)
        for l in lineas:
            if not (l.startswith("PROPUESTA") or l.startswith("Validado")):
                errores.append(f"{nombre}: {fid} con estado de contrato {l!r}")
        for c in cts:
            if not any(l.endswith(f"({c})") for l in lineas):
                errores.append(f"{nombre}: {fid} sin estado para {c}")
        archivos = RE_ARCHIVO.findall(codigo)
        for a in archivos:
            if not os.path.isfile(os.path.join(BACKEND_SRC, a)):
                errores.append(f"{nombre}: {fid} cita {a}, que no existe en backend/src")
        nombres = [os.path.basename(a) for a in archivos]
        if set(cts) & CON_PUERTO:
            n_codigo += 1
            if not any(n.startswith("Dummy") for n in nombres):
                errores.append(f"{nombre}: {fid} ({', '.join(cts)}) no cita un Dummy*.cs")
            if set(cts) & CON_HTTP and not any(n.startswith("Http") for n in nombres):
                errores.append(f"{nombre}: {fid} ({', '.join(cts)}) no cita un Http*.cs")
    if sin_tarea_f1:
        errores.append(f"{nombre}: {len(sin_tarea_f1)} ficha(s) de Fase 1 sin tarea en Pendientes_v4: "
                       f"{', '.join(sin_tarea_f1)}")
    return {"filas": len(cuerpo), "con_contrato": n_ct, "con_codigo": n_codigo, "contratos": sorted(citados)}


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    errores = []
    for p in (MD, DOCX, ESPEC_V4):
        if not os.path.exists(p):
            errores.append(f"no existe {p}")
    if errores:
        for e in errores:
            print("ERROR", e)
        return 1
    ids_spec = [f[0] for f in docx_utils.fichas(docx.Document(ESPEC_V4))]

    texto_md, filas_md = tabla_md()
    texto_dx, filas_dx, repite, horizontal = tabla_docx()
    r_md = verificar_filas("md", filas_md, ids_spec, errores)
    r_dx = verificar_filas("docx", filas_dx, ids_spec, errores)
    if len(filas_md) != len(filas_dx):
        errores.append(f"el md tiene {len(filas_md) - 1} filas y el docx {len(filas_dx) - 1}")
    if not repite:
        errores.append("docx: la fila de cabecera de la tabla no se repite (w:tblHeader)")
    if not horizontal:
        errores.append("docx: la sección no es horizontal")
    for nombre, texto in (("md", texto_md), ("docx", texto_dx)):
        hits = terminos.encontrar_prohibidos(texto)
        if hits:
            errores.append(f"{nombre}: {len(hits)} término(s) prohibido(s)")
    if "Fecha y commit de generación" not in texto_md or "Fecha y commit de generación" not in texto_dx:
        errores.append("falta la sección 'Fecha y commit de generación'")

    for nombre, r in (("md", r_md), ("docx", r_dx)):
        if r:
            print(f"{nombre}: {r['filas']} filas; {r['con_contrato']} con contrato; {r['con_codigo']} con puerto y "
                  f"código; contratos citados: {', '.join(r['contratos'])}")
    for e in errores:
        print("ERROR", e)
    if errores:
        print(f"{len(errores)} problema(s)")
        return 1
    print("matriz correcta")
    return 0


if __name__ == "__main__":
    sys.exit(main())
