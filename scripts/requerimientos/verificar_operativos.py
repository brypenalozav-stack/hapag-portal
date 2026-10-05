"""Fase 2: verifica los documentos operativos v4. Sale con 1 si algo falla.

- Pendientes_v4: IDs 1–118 conservados y al menos 118 filas; 0 celdas vacías en "Responsable / apoyo" y
  "Origen responsable"; el origen solo toma los 4 valores permitidos (ninguno "Gantt"); las tareas 37 y
  42–44 no contienen "sincroniz"; las tareas 102–106 están fuera del alcance (área Macros/RPX) y sin
  fechas; las filas con fecha son exactamente TAREAS_PLAN, con las fechas de config.CALENDARIO, y las
  demás dicen "Sin fecha en fuentes"; el zip contiene xl/charts/ y la hoja Resumen tiene un gráfico.
- NexusV2_v4: todas las funciones (las 16 filas del original) tienen R y A.
- Comparacion_v4: sin las hojas del sitio externo; 72 filas de Fase 1; 0 coincidencias de \\bPOC\\b y
  ninguna URL; cada evidencia file:line existe en el commit de línea base y su rango cabe en el archivo.
- No existe Gantt_Macros_Extraccion_QA_v4.xlsx.
- Los 3 archivos: 0 términos prohibidos en todas las celdas de todas las hojas.
"""
import datetime
import os
import re
import subprocess
import sys
import zipfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402
import editar_operativos_v4 as ed  # noqa: E402
import terminos  # noqa: E402
from linea_base import REPO  # noqa: E402

PENDIENTES = os.path.join(config.V4, "Portal_2_0_Pendientes_v4.xlsx")
NEXUS = os.path.join(config.V4, "Funcionalidades NexusV2_v4.xlsx")
COMPARACION = os.path.join(config.V4, "Comparacion_Requerimientos_Por_Fase_v4.xlsx")
GANTT_V4 = os.path.join(config.V4, "Gantt_Macros_Extraccion_QA_v4.xlsx")
RE_POC = re.compile(r"\bPOC\b")
RE_URL = re.compile(r"https?://|www\.", re.I)
RE_EVIDENCIA = re.compile(r"^([\w./-]+\.[\w]+|[\w./-]*Dockerfile):(\d+)(?:-(\d+))?$")


def celdas_texto(wb):
    for ws in wb.worksheets:
        yield ws.title, "título de hoja", ws.title
        for fila in ws.iter_rows():
            for c in fila:
                if isinstance(c.value, str):
                    yield ws.title, c.coordinate, c.value


def verificar_terminos(nombre, wb, errores):
    for hoja, coord, texto in celdas_texto(wb):
        if terminos.encontrar_prohibidos(texto):
            errores.append(f"{nombre}: término prohibido en {hoja}!{coord}")


def fecha(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    return v


def verificar_pendientes(errores):
    wb = openpyxl.load_workbook(PENDIENTES)
    ws = wb[ed.HOJA_PENDIENTES]
    cab = [c.value for c in ws[4]]
    col = {n: cab.index(n) for n in cab if n}
    for requerida in ("Responsable / apoyo", "Origen responsable", "Área responsable", "Ficha spec v4",
                      "Contrato de integración", "Fecha inicio", "Fecha objetivo", "Comentarios"):
        if requerida not in col:
            errores.append(f"Pendientes_v4: falta la columna {requerida!r}")
            return {}
    filas = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0] is not None]
    ids = [r[0] for r in filas]
    if ids[:118] != list(range(1, 119)):
        errores.append("Pendientes_v4: los IDs 1–118 no se conservan en orden")
    if len(filas) < 118:
        errores.append(f"Pendientes_v4: {len(filas)} filas (se esperaban al menos 118)")
    if len(set(ids)) != len(ids):
        errores.append("Pendientes_v4: IDs repetidos")

    por_origen = {}
    con_fecha = set()
    for r in filas:
        tid = r[0]
        resp, origen = r[col["Responsable / apoyo"]], r[col["Origen responsable"]]
        inicio, fin = fecha(r[col["Fecha inicio"]]), fecha(r[col["Fecha objetivo"]])
        coment = r[col["Comentarios"]] or ""
        if not resp:
            errores.append(f"Pendientes_v4: tarea {tid} sin 'Responsable / apoyo'")
        if not origen:
            errores.append(f"Pendientes_v4: tarea {tid} sin 'Origen responsable'")
        elif origen not in ed.ORIGENES or "gantt" in origen.lower():
            errores.append(f"Pendientes_v4: tarea {tid} con origen no permitido {origen!r}")
        por_origen[origen] = por_origen.get(origen, 0) + 1
        if tid == 37 or 42 <= tid <= 44:
            texto = " ".join(str(v) for v in r if v is not None).lower()
            if "sincroniz" in texto:
                errores.append(f"Pendientes_v4: la tarea {tid} contiene 'sincroniz'")
        if 102 <= tid <= 106:
            if ed.FUERA_ALCANCE_MACROS not in coment:
                errores.append(f"Pendientes_v4: la tarea {tid} no indica '{ed.FUERA_ALCANCE_MACROS}'")
            if inicio or fin:
                errores.append(f"Pendientes_v4: la tarea {tid} (Macros) tiene fechas")
        if inicio or fin:
            con_fecha.add(tid)
            if tid in ed.TAREAS_PLAN:
                esperado = ed.fechas_plan(tid)
                if (inicio, fin) != esperado:
                    errores.append(f"Pendientes_v4: tarea {tid} con fechas {inicio}–{fin} (config.CALENDARIO: "
                                   f"{esperado[0]}–{esperado[1]})")
                if origen != "Plan Portal 2.0":
                    errores.append(f"Pendientes_v4: tarea {tid} de TAREAS_PLAN con origen {origen!r}")
        elif ed.SIN_FECHA not in coment:
            errores.append(f"Pendientes_v4: tarea {tid} sin fecha y sin '{ed.SIN_FECHA}'")
        if tid not in ed.TAREAS_PLAN and origen == "Plan Portal 2.0":
            errores.append(f"Pendientes_v4: tarea {tid} con origen 'Plan Portal 2.0' fuera de TAREAS_PLAN")
        if str(r[2] or "").startswith("Validar contrato") and (inicio or fin or tid in ed.TAREAS_PLAN):
            errores.append(f"Pendientes_v4: la tarea {tid} 'Validar contrato' no debe tener fechas")
    if con_fecha != set(ed.TAREAS_PLAN):
        errores.append(f"Pendientes_v4: filas con fecha {sorted(con_fecha)} distintas de TAREAS_PLAN {sorted(ed.TAREAS_PLAN)}")

    with zipfile.ZipFile(PENDIENTES) as z:
        if not any(n.startswith("xl/charts/") for n in z.namelist()):
            errores.append("Pendientes_v4: el zip no contiene xl/charts/")
    if not wb["Resumen"]._charts:
        errores.append("Pendientes_v4: la hoja Resumen no tiene gráfico")
    verificar_terminos("Pendientes_v4", wb, errores)
    return {"total": len(filas), "nuevas": len(filas) - 118, "por_origen": por_origen,
            "con_fecha": sorted(con_fecha)}


def verificar_nexus(errores):
    original = openpyxl.load_workbook(ed.ORIG_NEXUS).active
    funciones_orig = [original.cell(r, 1).value for r in range(3, original.max_row + 1) if original.cell(r, 1).value]
    wb = openpyxl.load_workbook(NEXUS)
    ws = wb.active
    funciones = [(r, ws.cell(r, 1).value) for r in range(3, ws.max_row + 1) if ws.cell(r, 1).value]
    if [f for _, f in funciones] != funciones_orig:
        errores.append("NexusV2_v4: las funciones no coinciden con el original")
    completas = 0
    for r, nombre in funciones:
        if ws.cell(r, 5).value and ws.cell(r, 6).value:
            completas += 1
        else:
            errores.append(f"NexusV2_v4: la función {nombre!r} no tiene R y A")
    cab = [c.value for c in ws[2]]
    for requerida in ed.NEXUS_CABECERAS_NUEVAS:
        if requerida not in cab:
            errores.append(f"NexusV2_v4: falta la columna {requerida!r}")
    verificar_terminos("NexusV2_v4", wb, errores)
    return {"funciones": len(funciones), "con_R_y_A": completas}


def lineas_en_base(ruta, base, cache):
    if ruta not in cache:
        try:
            contenido = subprocess.check_output(["git", "-C", REPO, "show", f"{base}:{ruta}"], stderr=subprocess.DEVNULL)
            cache[ruta] = contenido.decode("utf-8", errors="replace").count("\n") + 1
        except subprocess.CalledProcessError:
            cache[ruta] = None
    return cache[ruta]


def verificar_comparacion(errores):
    wb = openpyxl.load_workbook(COMPARACION)
    for hoja in ed.HOJAS_ELIMINADAS:
        if hoja in wb.sheetnames:
            errores.append(f"Comparacion_v4: sigue la hoja {hoja!r}")
    ws = wb["Primera fase"]
    filas_f1 = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]]
    if len(filas_f1) != 72:
        errores.append(f"Comparacion_v4: Primera fase tiene {len(filas_f1)} filas (se esperaban 72)")
    for hoja, coord, texto in celdas_texto(wb):
        if RE_POC.search(texto):
            errores.append(f"Comparacion_v4: 'POC' en {hoja}!{coord}")
        if RE_URL.search(texto):
            errores.append(f"Comparacion_v4: URL en {hoja}!{coord}")

    base, _ = ed.commit_base()
    cache, n_evidencias = {}, 0
    for hoja, _, _ in ed.HOJAS_FASE:
        hs = wb[hoja]
        cab = [c.value for c in hs[4]]
        i_ev = cab.index("Evidencia (file:line)")
        i_cob = cab.index("Cobertura hapag-portal (C/P/N)")
        for r in hs.iter_rows(min_row=5, values_only=True):
            if not r[0]:
                continue
            if r[i_cob] not in ("C", "P", "N"):
                errores.append(f"Comparacion_v4: {r[0]} con cobertura {r[i_cob]!r}")
            if r[i_ev] in (None, "—"):
                if r[i_cob] in ("C", "P"):
                    errores.append(f"Comparacion_v4: {r[0]} ({r[i_cob]}) sin evidencia file:line")
                continue
            for ev in r[i_ev].split("; "):
                m = RE_EVIDENCIA.match(ev)
                if not m:
                    errores.append(f"Comparacion_v4: evidencia sin formato file:line en {r[0]}: {ev!r}")
                    continue
                n_evidencias += 1
                total = lineas_en_base(m.group(1), base, cache)
                fin = int(m.group(3) or m.group(2))
                if total is None:
                    errores.append(f"Comparacion_v4: {m.group(1)} no existe en el commit {base} ({r[0]})")
                elif fin > total:
                    errores.append(f"Comparacion_v4: {ev} excede las {total} líneas del archivo en {base} ({r[0]})")
    verificar_terminos("Comparacion_v4", wb, errores)
    return {"fase1": len(filas_f1), "evidencias": n_evidencias, "hojas": wb.sheetnames}


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    errores = []
    for p in (PENDIENTES, NEXUS, COMPARACION):
        if not os.path.exists(p):
            errores.append(f"no existe {p}")
    if os.path.exists(GANTT_V4):
        errores.append(f"existe {GANTT_V4}: el Gantt queda fuera del alcance")
    if errores:
        for e in errores:
            print("ERROR", e)
        return 1
    rp = verificar_pendientes(errores)
    rn = verificar_nexus(errores)
    rc = verificar_comparacion(errores)
    if rp:
        print(f"Pendientes_v4: {rp['total']} tareas ({rp['nuevas']} nuevas); por origen: "
              + ", ".join(f"{k} {v}" for k, v in rp["por_origen"].items()))
        print(f"  con fecha: {rp['con_fecha']}")
    print(f"NexusV2_v4: {rn['con_R_y_A']}/{rn['funciones']} funciones con R y A")
    print(f"Comparacion_v4: Primera fase {rc['fase1']} filas; {rc['evidencias']} evidencias file:line comprobadas; "
          f"hojas: {', '.join(rc['hojas'])}")
    for e in errores:
        print("ERROR", e)
    if errores:
        print(f"{len(errores)} problema(s)")
        return 1
    print("documentos operativos v4 correctos")
    return 0


if __name__ == "__main__":
    sys.exit(main())
