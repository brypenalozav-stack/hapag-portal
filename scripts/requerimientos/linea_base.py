"""Fase 0: congela la línea base de katu y exporta la matriz de 124 fichas.

- SHA-256 de los 6 originales -> docs/requerimientos/linea-base-katu.{json,md}
- katu\\v4\\Matriz_Fichas_Linea_Base_v4.xlsx con cobertura C/P/N y evidencia del research
  thoughts/shared/research/2026-10-05-actualizacion-requerimientos-portal-2-0.md (§2.1–§2.6).
"""
import datetime
import hashlib
import json
import os
import subprocess
import sys

import docx
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402
import docx_utils  # noqa: E402

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
DOCS = os.path.join(REPO, "docs", "requerimientos")

# Cobertura observada en el código (commit f49ea8a). Las fichas no listadas son "N" (No existe).
EVIDENCIA = {
    # M1
    "M1-01": ("P", "frontend/src/app/app.html:1-17; frontend/src/styles.scss:2-40"),
    "M1-02": ("P", "UsersController.cs:17-59 (global, sin ClientId); frontend/src/app/core/services/admin-user.service.ts:19-35"),
    "M1-04": ("P", "Client.cs:10 (país fijo); navbar.html:25-30 (solo lectura)"),
    "M1-05": ("P", "frontend/src/app/features/dashboard/dashboard.ts:30-47"),
    "M1-07": ("P", "RegisterCommandHandler.cs:20-117 (solo Client/CustomsAgent, sin aprobación)"),
    "M1-10": ("P", "ForgotPasswordCommandHandler.cs:14-43; ResetPasswordCommandHandler.cs:16-55; logout solo local auth.service.ts:81-87"),
    "M1-11": ("P", "ApplicationDbContext.cs:108-165; PermissionResolver.cs:10-31; HasPermissionAttribute.cs:5-52"),
    "M1-23": ("P", "AuditController.cs:15-23; SearchAuditQuery.cs:19-65 (sin escritor en ejecución)"),
    "M1-25": ("P", "NotificationsController.cs; NotificationPublisher.cs:17-51; navbar.html:33-40"),
    # M2
    "M2-06": ("P", "bl-list.ts:34-73; bl-detail.ts:108-142; GetMyBLsQueryHandler (solo ClientId propio)"),
    "M2-07": ("P", "BillOfLading.ShipmentType mostrado como columna Tipo (bl-list.html:56)"),
    "M2-09": ("P", "GetBLByNumberQueryHandler.cs:26 (sin TATC)"),
    # M3 / M4
    "M3-01": ("P", "GetLocalChargesByBLQueryHandler.cs:19-43 (lectura genérica)"),
    "M3-04": ("P", "CreateWarehouseChangeCommandHandler.cs:21-55 (sin rama gratuita)"),
    "M3-06": ("P", "WarehouseChangesController.cs:28-40 (sin historial de estados)"),
    "M3-12": ("P", "SubmitManifestAmendmentCommand (rectificación de manifiesto, sin cobro)"),
    "M3-13": ("P", "TransmitBLCommand.cs:50-59; regla BL_HOUSE_IN (sin cobro)"),
    "M3-14": ("P", "DeadlineCalculator.cs:9-36 (estado Overdue, sin cobro)"),
    "M4-01": ("P", "DemurrageExemption + Application/Admin/DemurrageExemptions (sin consumidor)"),
    "M4-03": ("P", "CreditClient + Application/Admin/CreditClients (sin consumidor)"),
    # M5 / M7
    "M5-01": ("P", "CreatePaymentCommand acepta Details; sin entidad ni vista de carro"),
    "M5-02": ("P", "CancelPaymentCommandHandler (cancela pagos no confirmados)"),
    "M5-03": ("P", "PaymentMethods.cs; GetPaymentMethodsQueryHandler.cs:15-31 (no consumido); payment-form.ts:44-58"),
    "M5-04": ("P", "Currency + GET config/currencies; CreatePaymentCommandHandler.cs:82-84 (moneda única)"),
    "M5-06": ("P", "Payment.DepositProofUrl (sin endpoint ni UI)"),
    "M5-07": ("P", "Método CreditLine sin validación contra CreditClient"),
    "M7-02": ("P", "GET payments/my; GET receipts/my; GetReceiptPdfQueryHandler.cs:31 (PDF placeholder)"),
    # M6 / M8 / M10
    "M6-08": ("P", "service-orders.ts:38-46 (tipos RELEASE, DECONSOLIDATION)"),
    "M8-05": ("P", "frontend/src/app/app.routes.ts:118-164 (rutas admin/*)"),
    "M8-06": ("P", "PermissionResolver.cs:10-20; LoginCommandHandler.cs:55 (rol ADMIN/USER)"),
    "M10-02": ("P", "FAQsController.cs; faq.ts:234-263 (FAQ por país, sin asistente)"),
    # NF
    "NF-01": ("P", "Webhooks y recibo idempotentes; creación de pago sin clave de idempotencia"),
    "NF-02": ("P", "PaymentStatus.cs:5-10 (sin historial de transiciones)"),
    "NF-04": ("P", "PaymentNumber único; ExternalReference (sin TransactionId persistido)"),
    "NF-05": ("P", "Filtro por ClientId en cada handler (GetBLByNumberQueryHandler.cs:26)"),
    "NF-07": ("P", "AesSecretProtector.cs; BCrypt; sin HTTPS/HSTS en Program.cs"),
    "NF-08": ("C", "Payment guarda solo referencia y resultado"),
    "NF-09": ("P", "SecretCredential + SecretResolver.cs:9-31 (tipos solo SII y Aduana)"),
    "NF-12": ("P", "CreatePaymentCommandHandler.cs:183-190 (Failed ante error; sin timeout)"),
    "NF-14": ("P", "AuditableEntityInterceptor.cs:27-55 (CreatedBy/ModifiedBy)"),
    "NF-15": ("P", "Sin valor anterior; AuditLog.OldValues sin escritor"),
    "NF-16": ("P", "AuditLog + borrado lógico; sin política de conservación"),
    "NF-19": ("P", "ImportBillsOfLadingCommandHandler.cs:25-119 (síncrono)"),
    "NF-21": ("P", "Grilla Bootstrap; un breakpoint (styles.scss:434)"),
    "NF-22": ("P", "UTC en auditoría; sin huso de referencia"),
    "NF-23": ("P", "Solo español (index.html:2 lang=es); sin LOCALE_ID"),
    "NF-24": ("P", "Staging en Railway (RAILWAY-DEPLOY.md)"),
    "NF-25": ("P", "Dockerfile + migraciones al arrancar; sin rollback"),
    "NF-26": ("P", "Solo /health (Program.cs:165)"),
    "NF-27": ("P", "ExceptionHandlingMiddleware + ILogger"),
}

# Fase indicada en el texto de la ficha cuando difiere del encabezado (defecto D7).
FASE_TEXTO = {
    "M3-02": "Texto: 'en su segunda fase'",
    "M6-02": "Texto: 'primera etapa sin pago'",
    "M3-17": "Nota propone Fase 0",
}

NOMBRES_MODULO = {
    "M1": "Acceso, usuarios y experiencia", "M2": "Disponibilidad de embarques",
    "M3": "Flujos de solicitud y servicios", "M4": "Reglas de negocio Nexus",
    "M5": "Carro y medios de pago", "M6": "Generación documental",
    "M7": "Facturación y estado de cuenta", "M8": "Administración interna",
    "M9": "Reportería", "M10": "Asistente virtual", "NF": "No funcionales",
}


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def git_head():
    try:
        return subprocess.check_output(["git", "-C", REPO, "rev-parse", "HEAD"], text=True).strip()
    except Exception:
        return "desconocido"


def escribir_linea_base():
    os.makedirs(DOCS, exist_ok=True)
    datos = {
        "generado": datetime.datetime.now().isoformat(timespec="seconds"),
        "commit_repo": git_head(),
        "carpeta": config.KATU,
        "archivos": [],
    }
    for nombre in config.ORIGINALES:
        p = os.path.join(config.KATU, nombre)
        st = os.stat(p)
        datos["archivos"].append({
            "nombre": nombre,
            "bytes": st.st_size,
            "modificado": datetime.datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"),
            "sha256": sha256(p),
            "tiene_v4": nombre in config.MAPA_V4,
        })
    with open(os.path.join(DOCS, "linea-base-katu.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump(datos, f, ensure_ascii=False, indent=2)
        f.write("\n")
    lineas = [
        "# Línea base de los documentos depurados (katu)",
        "",
        f"- Generado: {datos['generado']}",
        f"- Commit del repositorio: `{datos['commit_repo']}`",
        f"- Carpeta: `{config.KATU}`",
        "- Los originales no se modifican. Las versiones nuevas van a `katu\\v4\\` con sufijo `_v4`.",
        "- El Gantt de macros queda fuera del alcance: solo se conserva su hash.",
        "",
        "| Archivo | Bytes | Modificado | SHA-256 | Versión v4 |",
        "|---|---|---|---|---|",
    ]
    for a in datos["archivos"]:
        v4 = config.MAPA_V4.get(a["nombre"], "— (fuera de alcance o copia idéntica)")
        lineas.append(f"| {a['nombre']} | {a['bytes']} | {a['modificado']} | `{a['sha256']}` | {v4} |")
    with open(os.path.join(DOCS, "linea-base-katu.md"), "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lineas) + "\n")
    return datos


def exportar_matriz():
    d = docx.Document(config.ESPECIFICACION)
    filas = docx_utils.fichas(d)
    os.makedirs(config.V4, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fichas linea base"
    cab = ["ID", "Módulo", "Título", "Fase según encabezado", "Fase según texto",
           "Cobertura hapag-portal (C/P/N)", "Evidencia (file:line)"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0B2C5C")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for fid, titulo, fase in filas:
        mod = fid.split("-")[0]
        cob, ev = EVIDENCIA.get(fid, ("N", "—"))
        ws.append([fid, f"{mod} {NOMBRES_MODULO.get(mod, '')}".strip(), titulo, fase,
                   FASE_TEXTO.get(fid, "Igual al encabezado"), cob, ev])
    for col, ancho in zip("ABCDEFG", [9, 30, 55, 24, 26, 14, 80]):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    salida = config.ruta_v4("Matriz_Fichas_Linea_Base_v4.xlsx")
    wb.save(salida)
    desconocidas = sorted(set(EVIDENCIA) - {f[0] for f in filas})
    if desconocidas:
        raise SystemExit(f"Evidencia para fichas inexistentes: {desconocidas}")
    return salida, len(filas)


if __name__ == "__main__":
    datos = escribir_linea_base()
    salida, n = exportar_matriz()
    print(f"Línea base: {len(datos['archivos'])} archivos -> docs/requerimientos/linea-base-katu.json/.md")
    print(f"Matriz: {n} fichas -> {salida}")
