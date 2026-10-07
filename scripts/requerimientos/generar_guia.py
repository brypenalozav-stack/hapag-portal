"""Fase 3: Guía UI, accesibilidad e i18n (Portal 2.0 v4).

Genera, desde una sola estructura de bloques:
- docs/requerimientos/guia-ui-a11y-i18n.md
- katu\\v4\\Guia_UI_Accesibilidad_i18n_v4.docx (documento nuevo con python-docx)

Secciones: 1 tokens de diseño (valores de la Fase 5a §2 con contraste calculado según WCAG 2.2),
2 tipografía, 3 componentes y patrones, 4 checklist WCAG 2.2 AA (wcag22_a_aa.json, 55 criterios),
5 estrategia i18n con Transloco (Q8) y 6 glosario (docs/requerimientos/glosario-es-en.md).

Las utilidades de bloques (Markdown + Word) también las usa generar_matriz.py.
"""
import datetime
import json
import os
import re
import sys

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402
from linea_base import DOCS, REPO  # noqa: E402

AQUI = os.path.dirname(os.path.abspath(__file__))
WCAG_JSON = os.path.join(AQUI, "wcag22_a_aa.json")
GLOSARIO_MD = os.path.join(DOCS, "glosario-es-en.md")
SALIDA_MD = os.path.join(DOCS, "guia-ui-a11y-i18n.md")
NOMBRE_DOCX = "Guia_UI_Accesibilidad_i18n_v4.docx"

TITULO = "Guía de UI, accesibilidad e idiomas — Portal 2.0 v4"
SECCIONES = [
    "1. Tokens de diseño",
    "2. Tipografía",
    "3. Componentes y patrones",
    "4. Checklist WCAG 2.2 AA",
    "5. Estrategia i18n con Transloco",
    "6. Glosario ES/EN",
]
METODOS = ("eslint", "axe", "teclado", "NVDA", "Lighthouse", "revisión visual")

AZUL_HL = "0B2C5C"
CEBRA = "F5F7FA"


# ================================================================ contraste (WCAG 2.2, luminancia relativa)

def _lineal(c):
    c = c / 255
    return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4


def luminancia(hex_color):
    h = hex_color.lstrip("#")
    r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * _lineal(r) + 0.7152 * _lineal(g) + 0.0722 * _lineal(b)


def contraste(a, b):
    la, lb = sorted((luminancia(a), luminancia(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def fmt_ratio(x):
    return f"{x:.2f}".replace(".", ",") + ":1"


# ================================================================ bloques -> Markdown / Word

class Documento:
    """Secuencia de bloques: ('h', nivel, texto), ('p', texto), ('lista', [..]), ('codigo', texto),
    ('tabla', cabeceras, filas, anchos_cm). Las celdas pueden ser str o list[str] (varias líneas)."""

    def __init__(self, titulo, subtitulo=None):
        self.titulo = titulo
        self.subtitulo = subtitulo
        self.bloques = []

    def h(self, nivel, texto):
        self.bloques.append(("h", nivel, texto))

    def p(self, texto):
        self.bloques.append(("p", texto))

    def lista(self, items):
        self.bloques.append(("lista", list(items)))

    def codigo(self, texto):
        self.bloques.append(("codigo", texto))

    def tabla(self, cabeceras, filas, anchos_cm=None):
        self.bloques.append(("tabla", list(cabeceras), [list(f) for f in filas], anchos_cm))


def _celda_md(v):
    if isinstance(v, (list, tuple)):
        return "<br>".join(_celda_md(x) for x in v)
    return str(v).replace("|", "\\|").replace("\n", "<br>")


def a_markdown(d):
    out = [f"# {d.titulo}", ""]
    if d.subtitulo:
        out += [d.subtitulo, ""]
    for b in d.bloques:
        if b[0] == "h":
            out += ["#" * (b[1] + 1) + " " + b[2], ""]
        elif b[0] == "p":
            out += [b[1], ""]
        elif b[0] == "lista":
            out += [f"- {x}" for x in b[1]] + [""]
        elif b[0] == "codigo":
            out += ["```", b[1].rstrip("\n"), "```", ""]
        elif b[0] == "tabla":
            _, cab, filas, _ = b
            out.append("| " + " | ".join(_celda_md(c) for c in cab) + " |")
            out.append("|" + "---|" * len(cab))
            for f in filas:
                out.append("| " + " | ".join(_celda_md(c) for c in f) + " |")
            out.append("")
    return "\n".join(out).rstrip("\n") + "\n"


RE_INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")


def _runs(parrafo, texto, tam=None, negrita=False, color=None):
    for parte in RE_INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith("`") and parte.endswith("`"):
            r = parrafo.add_run(parte[1:-1])
            r.font.name = "Consolas"
            _fuente_rfonts(r._r, "Consolas")
        elif parte.startswith("**") and parte.endswith("**"):
            r = parrafo.add_run(parte[2:-2])
            r.bold = True
        else:
            r = parrafo.add_run(parte)
        if negrita:
            r.bold = True
        if tam:
            r.font.size = Pt(tam)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def _fuente_rfonts(el, nombre):
    rpr = el.get_or_add_rPr() if hasattr(el, "get_or_add_rPr") else el
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), nombre)


def _estilo(doc, nombre, tam, color=None, negrita=None):
    s = doc.styles[nombre]
    s.font.name = "Calibri"
    s.font.size = Pt(tam)
    _fuente_rfonts(s.element.get_or_add_rPr(), "Calibri")
    if color:
        s.font.color.rgb = RGBColor.from_string(color)
    if negrita is not None:
        s.font.bold = negrita
    return s


def _sombrear(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def _fila_propiedad(fila, tag):
    trpr = fila._tr.get_or_add_trPr()
    el = OxmlElement(tag)
    el.set(qn("w:val"), "true")
    trpr.append(el)


def _texto_celda(celda, valor, tam, negrita=False, color=None):
    lineas = valor if isinstance(valor, (list, tuple)) else str(valor).split("\n")
    p = celda.paragraphs[0]
    for i, linea in enumerate(lineas):
        if i:
            p = celda.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        _runs(p, str(linea), tam=tam, negrita=negrita, color=color)


def _numero_pagina(parrafo):
    for tipo, texto in (("begin", None), (None, " PAGE "), ("end", None)):
        r = OxmlElement("w:r")
        if tipo:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), tipo)
            r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = texto
            r.append(it)
        parrafo._p.append(r)


def a_docx(d, ruta, horizontal=False, tam_tabla=9, pie=None, portada=None):
    doc = docx.Document()
    _estilo(doc, "Normal", 10)
    _estilo(doc, "Title", 24, AZUL_HL)
    _estilo(doc, "Subtitle", 13, "33424F")
    _estilo(doc, "Heading 1", 16, AZUL_HL, True)
    _estilo(doc, "Heading 2", 13, AZUL_HL, True)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    if horizontal:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    margen = Cm(1.5) if horizontal else Cm(2.0)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = margen
    ancho_util = (sec.page_width - sec.left_margin - sec.right_margin) / 360000  # cm

    props = doc.core_properties
    props.title = d.titulo
    props.author = config.EJECUTA
    props.language = "es-CL"

    if pie:
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _runs(fp, pie + " · página ", tam=8, color="4A5568")
        _numero_pagina(fp)

    doc.add_paragraph(d.titulo, style="Title")
    if d.subtitulo:
        doc.add_paragraph(d.subtitulo, style="Subtitle")
    for linea in portada or []:
        _runs(doc.add_paragraph(), linea)

    for b in d.bloques:
        if b[0] == "h":
            doc.add_paragraph(b[2], style=f"Heading {b[1]}")
        elif b[0] == "p":
            _runs(doc.add_paragraph(), b[1])
        elif b[0] == "lista":
            for x in b[1]:
                _runs(doc.add_paragraph(style="List Bullet"), x)
        elif b[0] == "codigo":
            for linea in b[1].rstrip("\n").split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.5)
                r = p.add_run(linea)
                r.font.size = Pt(9)
                _fuente_rfonts(r._r, "Consolas")
            doc.add_paragraph()
        elif b[0] == "tabla":
            _, cab, filas, anchos = b
            if anchos is None:
                anchos = [ancho_util / len(cab)] * len(cab)
            else:
                escala = ancho_util / sum(anchos)
                anchos = [a * escala for a in anchos]
            t = doc.add_table(rows=1, cols=len(cab))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = False
            enc = t.rows[0]
            _fila_propiedad(enc, "w:tblHeader")
            _fila_propiedad(enc, "w:cantSplit")
            for i, c in enumerate(cab):
                _texto_celda(enc.cells[i], c, tam_tabla, negrita=True, color="FFFFFF")
                _sombrear(enc.cells[i], AZUL_HL)
            for n, f in enumerate(filas):
                fila = t.add_row()
                _fila_propiedad(fila, "w:cantSplit")
                celdas = fila.cells
                for i, v in enumerate(f):
                    _texto_celda(celdas[i], v, tam_tabla)
                    if n % 2:
                        _sombrear(celdas[i], CEBRA)
            for fila in t.rows:
                for i, c in enumerate(fila.cells):
                    c.width = Cm(anchos[i])
            doc.add_paragraph()
    doc.save(ruta)
    return ruta


def escribir_md(ruta, texto):
    with open(ruta, "w", encoding="utf-8", newline="\n") as f:
        f.write(texto)


# ================================================================ 1. tokens (Fase 5a §2, DC4)

# (token Sass, valor, uso, variante oscura). Lista cerrada de la Fase 5a §2 del plan.
TOKENS = [
    ("$hl-dark", "#33424f", "Secundario (`$secondary`), barra de navegación y menú lateral",
     "Sin cambio (#33424f); también es la superficie de tarjetas en oscuro"),
    ("$hl-orange", "#ff6600", "Solo logotipo y acentos no textuales", "Sin cambio (#ff6600)"),
    ("$hl-orange-700", "#b84a00", "Primario (`$primary`): botones, enlaces y borde de foco de campos",
     "#b84a00 solo como fondo de botón con texto blanco; los enlaces usan `--hl-dark-text` subrayado"),
    ("$hl-green", "#009840", "Acentos no textuales de éxito (íconos grandes, barras)", "Sin cambio (#009840)"),
    ("$hl-green-700", "#007a33", "Éxito (`$success`): texto e insignias",
     "#007a33 como fondo de insignia con texto blanco"),
    ("$hl-blue", "#004d6c", "Información (`$info`) y anillo de foco sobre fondos claros",
     "#004d6c como fondo; el foco pasa a `--hl-white`"),
    ("$hl-light-gray", "#f5f7fa", "Fondo de página (`$body-bg`)", "#1a202c (`$hl-dark-bg`)"),
    ("$hl-border", "#e2e8f0", "Bordes de tarjetas, tablas y separadores", "#4a5568 (`$hl-text-muted`)"),
    ("$hl-body-color", "#2d3748", "Texto de cuerpo (`$body-color`)", "#e2e8f0 (`$hl-dark-text`)"),
    ("$hl-white", "#ffffff", "Superficies (tarjetas, tablas) y texto sobre fondos oscuros",
     "Superficie: #33424f (`$hl-dark`); texto y foco: #ffffff"),
    ("$hl-dark-bg", "#1a202c", "Fondo de página en tema oscuro", "—"),
    ("$hl-dark-text", "#e2e8f0", "Texto en tema oscuro", "—"),
    ("$hl-surface-muted", "#f8fafc", "Superficies secundarias (cabeceras de tabla, filtros)", "#33424f (`$hl-dark`)"),
    ("$hl-text-muted", "#4a5568", "Texto secundario; reemplaza a #718096 (4,02:1 sobre blanco)",
     "#e2e8f0 (`$hl-dark-text`)"),
]
# Variable CSS sin token Sass propio: --hl-focus = $hl-blue (en .hl-navbar y .hl-sidebar, $hl-white).
FOCO = ("--hl-focus", "#004d6c", "Anillo de foco `:focus-visible` (3 px, desplazamiento 2 px). En `.hl-navbar` y "
        "`.hl-sidebar` vale #ffffff", "#ffffff (`$hl-white`)")

# Contraste esperado según el plan (Key Discoveries y DC4). generar_guia aborta si el cálculo no coincide.
ESPERADO_PLAN = {
    ("#ff6600", "#ffffff"): 2.94, ("#b84a00", "#ffffff"): 5.23, ("#b84a00", "#f5f7fa"): 4.87,
    ("#009840", "#ffffff"): 3.77, ("#007a33", "#ffffff"): 5.48, ("#004d6c", "#ffffff"): 9.22,
    ("#33424f", "#ffffff"): 10.33, ("#4a5568", "#ffffff"): 7.53,
}

# (tema, primer plano, fondo, uso). Los tokens se resuelven con TOKENS.
PARES = [
    ("Claro", "$hl-body-color", "$hl-white", "Texto de cuerpo sobre tarjetas"),
    ("Claro", "$hl-body-color", "$hl-light-gray", "Texto de cuerpo sobre el fondo de página"),
    ("Claro", "$hl-text-muted", "$hl-white", "Texto secundario"),
    ("Claro", "$hl-text-muted", "$hl-light-gray", "Texto secundario sobre el fondo de página"),
    ("Claro", "$hl-text-muted", "$hl-surface-muted", "Texto secundario en cabeceras de tabla y filtros"),
    ("Claro", "$hl-orange-700", "$hl-white", "Enlaces y botones de contorno"),
    ("Claro", "$hl-orange-700", "$hl-light-gray", "Enlaces sobre el fondo de página"),
    ("Claro", "$hl-white", "$hl-orange-700", "Texto de botón primario"),
    ("Claro", "$hl-green-700", "$hl-white", "Texto de éxito"),
    ("Claro", "$hl-green-700", "$hl-light-gray", "Texto de éxito sobre el fondo de página"),
    ("Claro", "$hl-white", "$hl-green-700", "Insignia de éxito"),
    ("Claro", "$hl-blue", "$hl-white", "Texto informativo y anillo de foco"),
    ("Claro", "$hl-blue", "$hl-light-gray", "Anillo de foco sobre el fondo de página"),
    ("Claro", "$hl-dark", "$hl-white", "Títulos y elementos secundarios"),
    ("Claro", "$hl-white", "$hl-dark", "Texto y foco en la barra de navegación y el menú lateral"),
    ("Claro", "$hl-orange", "$hl-white", "Logotipo y acentos: no apto para texto ni para bordes de controles"),
    ("Claro", "$hl-orange", "$hl-dark", "Logotipo sobre la barra de navegación"),
    ("Claro", "$hl-green", "$hl-white", "Acento no textual de éxito; no apto para texto"),
    ("Claro", "$hl-border", "$hl-white", "Separadores decorativos (ver nota sobre campos de formulario)"),
    ("Oscuro", "$hl-dark-text", "$hl-dark-bg", "Texto de cuerpo"),
    ("Oscuro", "$hl-dark-text", "$hl-dark", "Texto sobre superficies"),
    ("Oscuro", "$hl-white", "$hl-dark-bg", "Anillo de foco"),
    ("Oscuro", "$hl-white", "$hl-orange-700", "Texto de botón primario"),
    ("Oscuro", "$hl-orange-700", "$hl-dark-bg", "Enlace naranja sobre fondo oscuro: no permitido como texto"),
    ("Oscuro", "$hl-text-muted", "$hl-dark-bg", "Separadores decorativos"),
]


def valor_token(nombre):
    for t, v, _, _ in TOKENS:
        if t == nombre:
            return v
    raise KeyError(nombre)


def filas_contraste():
    for (fg, bg), esperado in ESPERADO_PLAN.items():
        calculado = round(contraste(fg, bg), 2)
        if calculado != esperado:
            raise SystemExit(f"Contraste {fg} sobre {bg}: calculado {calculado}, el plan indica {esperado}")
    filas = []
    for tema, fg, bg, uso in PARES:
        x = contraste(valor_token(fg), valor_token(bg))
        filas.append([tema, f"`{fg}` {valor_token(fg)}", f"`{bg}` {valor_token(bg)}", fmt_ratio(x),
                      "Cumple" if x >= 4.5 else "No cumple", "Cumple" if x >= 3 else "No cumple", uso])
    return filas


def seccion_tokens(d):
    d.h(1, SECCIONES[0])
    d.p("Los valores son los que fija la Fase 5a §2 del plan (decisión DC4). En el código viven solo en "
        "`frontend/src/styles/_tokens.scss`, como variables Sass sin CSS emitido; el resto del frontend usa "
        "`var(--hl-…)` o las variables de Bootstrap configuradas con `@use 'bootstrap/scss/bootstrap' with (…)`. "
        "Hoy la interfaz muestra el azul por defecto de Bootstrap (`--bs-primary: #0d6efd`) porque los overrides de "
        "`styles.scss` no se aplican: el cambio visual de la Fase 5a es mayor de lo que sugiere el código actual.")
    d.h(2, "1.1 Lista cerrada de tokens")
    d.p("Es la misma lista que usa el prototipo de la Fase 4. Un color nuevo se agrega aquí, en `_tokens.scss` y en el "
        "prototipo en el mismo cambio; ningún otro archivo del frontend contiene valores hexadecimales "
        "(`npm run check:hex`).")
    filas = [[f"`{t}`", f"`--{t[1:]}`", v, uso, oscuro] for t, v, uso, oscuro in TOKENS]
    filas.append(["—", f"`{FOCO[0]}`", FOCO[1], FOCO[2], FOCO[3]])
    d.tabla(["Token Sass", "Variable CSS", "Valor", "Uso", "Variante oscura"], filas, [3.2, 3.6, 1.8, 5.2, 4.6])
    d.h(2, "1.2 Contraste de los pares documentados")
    d.p("Contraste calculado por `generar_guia.py` con la fórmula de luminancia relativa de WCAG 2.2. Umbrales: "
        "4,5:1 para texto normal (1.4.3) y 3:1 para texto grande, componentes de interfaz y foco (1.4.3 y 1.4.11).")
    d.tabla(["Tema", "Primer plano", "Fondo", "Contraste", "Texto normal (4,5:1)", "Texto grande y no textual (3:1)",
             "Uso"], filas_contraste(), [1.4, 3.2, 3.2, 1.8, 2.0, 2.4, 4.4])
    d.h(2, "1.3 Reglas de uso del color")
    d.lista([
        "`#ff6600` (`$hl-orange`) se reserva para el logotipo, que está exento de 1.4.3, y para acentos no textuales "
        "que no transmiten información por sí solos. Con 2,94:1 sobre blanco no alcanza ni el 3:1 de 1.4.11, por eso "
        "no se usa en texto, bordes de controles, íconos de estado ni foco.",
        "El primario es `#b84a00` (`$hl-orange-700`): 5,23:1 sobre blanco y 4,87:1 sobre `#f5f7fa`.",
        "El éxito en texto es `#007a33` (`$hl-green-700`, 5,48:1). `#009840` (3,77:1) solo se usa en acentos no "
        "textuales.",
        "El texto secundario es `#4a5568` (`$hl-text-muted`, 7,53:1); `#718096` (4,02:1) no se usa.",
        "El color nunca es el único medio para transmitir un estado (1.4.1): las insignias de estado llevan texto y "
        "los errores llevan ícono y mensaje.",
        "Campos de formulario: `$hl-border` da 1,23:1 sobre blanco. Si el borde es lo único que identifica el campo, "
        "1.4.11 exige 3:1. Observación para la Fase 5a: antes de cerrar M11-04, el borde de los campos debe usar un "
        "tono con al menos 3:1 (por ejemplo `$hl-text-muted`, 7,53:1) o el campo debe tener otro indicador visual "
        "con ese contraste. En tarjetas, tablas y separadores el borde es decorativo y no requiere 3:1.",
        "Tema oscuro: los tokens oscuros se declaran en `[data-bs-theme=\"dark\"]`; el selector de tema queda "
        "pendiente (tarea M11-07). Sobre `#1a202c` el naranja `#b84a00` da 3,12:1, por eso en oscuro los enlaces "
        "usan `--hl-dark-text` subrayado y el naranja queda como fondo de botón con texto blanco.",
    ])


# ================================================================ 2. tipografía

def seccion_tipografia(d):
    d.h(1, SECCIONES[1])
    d.p("Ambas familias tienen licencia SIL Open Font License 1.1 (OFL-1.1), que permite incrustarlas y servirlas "
        "desde el portal sin costo. Se configuran en la Fase 5a mediante `$font-family-sans-serif` y "
        "`$headings-font-family` de Bootstrap.")
    d.tabla(["Uso", "Familia y respaldo", "Peso", "Tamaño", "Licencia"], [
        ["Cuerpo, formularios y tablas", "Inter, system-ui, -apple-system, sans-serif", "400; 600 para énfasis",
         "1rem (16 px) con interlineado 1,5", "OFL-1.1"],
        ["Títulos (h1–h6)", "Montserrat, Inter, system-ui, sans-serif", "600 (`$headings-font-weight`)",
         "Escala de Bootstrap en rem", "OFL-1.1"],
        ["Montos, números de BL y fechas en tablas", "Inter con `font-variant-numeric: tabular-nums`", "400",
         "0,875rem como mínimo", "OFL-1.1"],
    ], [3.6, 4.6, 3.0, 3.6, 1.6])
    d.lista([
        "Todos los tamaños se expresan en `rem`, para que el texto crezca al 200 % sin pérdida de contenido (1.4.4).",
        "Ningún texto informativo baja de 0,75rem (12 px); el texto de cuerpo no baja de 0,875rem (14 px).",
        "Sin alturas fijas en contenedores de texto: el contenido admite el espaciado de 1.4.12 (interlineado 1,5; "
        "párrafos 2; letras 0,12; palabras 0,16 veces el tamaño de la fuente).",
        "No se usan imágenes de texto (1.4.5); el logotipo es la única excepción.",
        "Los títulos siguen una jerarquía sin saltos (h1 único por pantalla, luego h2, h3).",
    ])


# ================================================================ 3. componentes y patrones

def seccion_componentes(d):
    d.h(1, SECCIONES[2])
    d.p("Patrones obligatorios para las pantallas de Angular y para el prototipo. Todos los textos, incluidos "
        "`aria-label`, `alt`, `title` y `placeholder`, salen de claves Transloco (sección 5).")

    d.h(2, "3.1 Tablas de datos")
    d.lista([
        "Cada tabla lleva `<caption>`; si el título ya es visible, el caption puede ir con `class=\"visually-hidden\"`.",
        "Los `th` de cabecera llevan `scope=\"col\"`; los encabezados de fila, `scope=\"row\"`.",
        "Las tablas no se usan para maquetar.",
        "El contenedor con desplazamiento horizontal (`.table-responsive`) es enfocable (`tabindex=\"0\"`) y tiene "
        "nombre accesible, para que se pueda recorrer con teclado.",
        "El orden de columnas se anuncia con `aria-sort` en el `th` activo.",
        "Los montos se alinean a la derecha y siempre muestran el código ISO de la moneda.",
    ])
    d.codigo('<table class="table">\n'
             '  <caption class="visually-hidden">{{ \'bl.list.caption\' | transloco }}</caption>\n'
             '  <thead>\n'
             '    <tr>\n'
             '      <th scope="col">{{ \'bl.list.number\' | transloco }}</th>\n'
             '      <th scope="col">{{ \'bl.list.vessel\' | transloco }}</th>\n'
             '    </tr>\n'
             '  </thead>\n'
             '</table>')

    d.h(2, "3.2 Formularios")
    d.lista([
        "Cada control tiene `label for` asociado a su `id` (regla `label-has-associated-control` de angular-eslint).",
        "Las ayudas y el mensaje de error se vinculan con `aria-describedby`; el campo con error lleva "
        "`aria-invalid=\"true\"`.",
        "Los campos obligatorios usan `required` y lo indican en texto, no solo con un asterisco de color.",
        "Los campos de datos personales llevan `autocomplete` (`email`, `username`, `current-password`, "
        "`new-password`, `organization`, `tel`) (1.3.5).",
        "Al enviar con errores se muestra un **resumen de errores** al inicio del formulario: contenedor con "
        "`tabindex=\"-1\"` que recibe el foco, un título y un enlace por error que lleva al campo.",
        "El mensaje de error dice qué pasó y cómo corregirlo (3.3.3), por ejemplo el formato esperado del RUT o NIT.",
        "Los datos ya ingresados en el mismo proceso no se vuelven a pedir (3.3.7).",
        "El pago muestra un resumen para confirmar antes de enviar (3.3.4).",
    ])
    d.codigo('<label for="rut" class="form-label">{{ \'register.form.taxId\' | transloco }}</label>\n'
             '<input id="rut" class="form-control" required autocomplete="off"\n'
             '       aria-describedby="rut-ayuda rut-error" [attr.aria-invalid]="rutInvalido">\n'
             '<div id="rut-ayuda" class="form-text">{{ \'register.form.taxIdHelp\' | transloco }}</div>\n'
             '<div id="rut-error" class="invalid-feedback">{{ \'register.form.taxIdError\' | transloco }}</div>')

    d.h(2, "3.3 Mensajes dinámicos con aria-live")
    d.lista([
        "El shell tiene dos regiones `visually-hidden`: `aria-live=\"polite\"` y `aria-live=\"assertive\"`. Las "
        "escribe `LiveAnnouncerService` (Fase 5c), sin CDK.",
        "**Pagos:** el resultado (confirmado, rechazado, pendiente) se anuncia en la región polite sin mover el foco.",
        "**Cargas:** la importación de BL anuncia el inicio y el resultado con el número de registros.",
        "**Errores de envío:** se anuncian en la región assertive y se muestra el resumen de errores.",
        "Los avisos temporales (toasts) no desaparecen antes de 5 segundos y se pueden cerrar con teclado.",
    ])

    d.h(2, "3.4 Estados vacío y error (NF-11)")
    d.p("Componente `state-message` (Fase 5c) con `kind: 'empty' | 'error'`. Se usa en dashboard, listado y detalle "
        "de BL, listado de pagos y comprobantes.")
    d.tabla(["Estado", "Condición", "Rol", "Texto (ES)", "Acción"], [
        ["Vacío (`empty`)", "HTTP 200 con lista vacía", "`role=\"status\"`", "No hay datos", "Ninguna"],
        ["Error (`error`)", "HTTP 5xx o estado 0 (sin conexión)", "`role=\"alert\"`",
         "Servicio temporalmente no disponible", "Botón Reintentar, que repite la consulta"],
    ], [2.6, 3.8, 2.6, 4.2, 4.0])

    d.h(2, "3.5 Navegación, foco y movimiento")
    d.lista([
        "Enlace para saltar al contenido (`.hl-skip-link`) como primer elemento enfocable; lleva a "
        "`<main id=\"contenido-principal\" tabindex=\"-1\">`.",
        "Regiones `<header>`, `<nav>` con `aria-label` y `<main>`; el enlace activo del menú lleva "
        "`aria-current=\"page\"`.",
        "Foco visible en todo elemento interactivo: `:focus-visible { outline: 3px solid var(--hl-focus); "
        "outline-offset: 2px; }`. Ningún elemento fijo tapa el foco (2.4.11).",
        "Botón de menú con `type=\"button\"`, nombre accesible, `aria-controls` y `aria-expanded`.",
        "Selector de idioma ES/EN como grupo de botones con `aria-pressed`.",
        "Los botones de solo ícono miden al menos 24×24 px (2.5.8) y llevan `aria-label`; los íconos decorativos, "
        "`aria-hidden=\"true\"`.",
        "Con `prefers-reduced-motion: reduce`, animaciones y transiciones bajan a 0,01 ms y "
        "`scroll-behavior` pasa a `auto`.",
    ])


# ================================================================ 4. checklist WCAG 2.2 AA

SIN_MULTIMEDIA = ("No aplica: el portal no publica audio ni vídeo", "revisión visual",
                  "Todas (se revisa al incorporar contenido multimedia)")

# id -> (aplica, métodos, pantallas)
CHECKLIST = {
    "1.1.1": ("Sí", "eslint, axe, NVDA", "Todas: logotipo, íconos y gráficos del dashboard"),
    "1.2.1": SIN_MULTIMEDIA, "1.2.2": SIN_MULTIMEDIA, "1.2.3": SIN_MULTIMEDIA, "1.2.4": SIN_MULTIMEDIA,
    "1.2.5": SIN_MULTIMEDIA,
    "1.3.1": ("Sí", "eslint, axe, NVDA", "Listado de BL, detalle de BL, pagos, comprobantes y formularios"),
    "1.3.2": ("Sí", "NVDA, revisión visual", "Todas"),
    "1.3.3": ("Sí", "revisión visual", "Todas (instrucciones sin depender de forma, color o posición)"),
    "1.3.4": ("Sí", "revisión visual", "Todas, en móvil vertical y horizontal (NF-21)"),
    "1.3.5": ("Sí", "axe, revisión visual", "Inicio de sesión, registro y recuperación de contraseña"),
    "1.4.1": ("Sí", "axe, revisión visual", "Insignias de estado, errores de formulario y enlaces en texto"),
    "1.4.2": ("No aplica: el portal no reproduce audio", "revisión visual", "Todas"),
    "1.4.3": ("Sí", "axe, Lighthouse", "Todas (tokens de la sección 1)"),
    "1.4.4": ("Sí", "revisión visual", "Todas, con zoom del navegador al 200 %"),
    "1.4.5": ("Sí", "revisión visual", "Todas (logotipo exento)"),
    "1.4.10": ("Sí", "revisión visual", "Todas, a 320 px de ancho (NF-21)"),
    "1.4.11": ("Sí", "revisión visual", "Campos de formulario, foco, íconos de estado y botones de contorno"),
    "1.4.12": ("Sí", "revisión visual", "Todas, con un marcador de espaciado de texto"),
    "1.4.13": ("Sí", "teclado, revisión visual", "Tooltips y menús desplegables"),
    "2.1.1": ("Sí", "eslint, teclado", "Todas"),
    "2.1.2": ("Sí", "teclado", "Modales, menú lateral móvil y selectores de fecha"),
    "2.1.4": ("Sí: no se definen atajos de un solo carácter", "teclado", "Todas"),
    "2.2.1": ("Sí", "teclado, revisión visual", "Expiración de sesión y formulario de pago"),
    "2.2.2": ("Sí", "revisión visual", "Dashboard (comunicados) e indicadores de carga"),
    "2.3.1": ("Sí", "revisión visual", "Todas"),
    "2.4.1": ("Sí", "axe, teclado", "Todas (enlace para saltar al contenido y regiones)"),
    "2.4.2": ("Sí", "axe, Lighthouse", "Todas (título traducido por ruta)"),
    "2.4.3": ("Sí", "teclado, NVDA", "Todas"),
    "2.4.4": ("Sí", "axe, NVDA", "Todas"),
    "2.4.5": ("Sí", "revisión visual", "Menú lateral y búsqueda de BL"),
    "2.4.6": ("Sí", "NVDA, revisión visual", "Todas"),
    "2.4.7": ("Sí", "teclado", "Todas (`:focus-visible` con `--hl-focus`)"),
    "2.4.11": ("Sí", "teclado", "Todas (barra superior fija y paneles laterales)"),
    "2.5.1": ("Sí", "revisión visual", "Todas (sin gestos multipunto ni de trayectoria)"),
    "2.5.2": ("Sí", "revisión visual", "Botones y enlaces (la acción ocurre al soltar)"),
    "2.5.3": ("Sí", "axe, NVDA", "Botones e íconos con texto visible"),
    "2.5.4": ("No aplica: no hay funciones por movimiento del dispositivo", "revisión visual", "Todas"),
    "2.5.7": ("Sí", "teclado, revisión visual", "Importación de BL y adjuntos (alternativa con botón)"),
    "2.5.8": ("Sí", "axe, revisión visual", "Botones de ícono, paginación y selector de idioma"),
    "3.1.1": ("Sí", "axe, Lighthouse", "Todas (`html[lang]` según el idioma activo)"),
    "3.1.2": ("Sí", "NVDA, revisión visual", "Selector de idioma (opción en el otro idioma con su `lang`)"),
    "3.2.1": ("Sí", "teclado", "Todas"),
    "3.2.2": ("Sí", "teclado, NVDA", "Selectores de país e idioma y filtros de listados"),
    "3.2.3": ("Sí", "revisión visual", "Todas (shell común)"),
    "3.2.4": ("Sí", "revisión visual", "Todas (términos del glosario de la sección 6)"),
    "3.2.6": ("Sí", "revisión visual", "Preguntas frecuentes y asistente en la misma posición"),
    "3.3.1": ("Sí", "axe, NVDA", "Formularios (`aria-invalid` y resumen de errores)"),
    "3.3.2": ("Sí", "eslint, axe", "Formularios"),
    "3.3.3": ("Sí", "NVDA, revisión visual", "Formularios (formato de RUT o NIT, fechas y montos)"),
    "3.3.4": ("Sí", "teclado, revisión visual", "Formulario de pago y carro por moneda (confirmación)"),
    "3.3.7": ("Sí", "revisión visual", "Registro y pago"),
    "3.3.8": ("Sí", "teclado, revisión visual", "Inicio de sesión (pegar contraseña, gestor de contraseñas)"),
    "4.1.2": ("Sí", "eslint, axe, NVDA", "Todas (menú con `aria-expanded`, selector ES/EN con `aria-pressed`)"),
    "4.1.3": ("Sí", "NVDA", "Pago, importación de BL y estados de NF-11"),
}


def cargar_wcag():
    with open(WCAG_JSON, encoding="utf-8") as f:
        return json.load(f)["criterios"]


def seccion_checklist(d):
    d.h(1, SECCIONES[3])
    d.p("Nivel de conformidad decidido en Q9: los 55 criterios A y AA de la Recomendación W3C *Web Content "
        "Accessibility Guidelines (WCAG) 2.2* (https://www.w3.org/TR/WCAG22/); 4.1.1 queda fuera por obsoleto. "
        "La lista canónica está en `scripts/requerimientos/wcag22_a_aa.json`.")
    d.p("Evidencia exigida por Q9: `ng lint` sin errores en CI; axe con 0 violaciones en ES y EN; Lighthouse "
        "Accesibilidad ≥ 95; prueba manual con teclado y NVDA siguiendo esta tabla.")
    d.lista([
        "**eslint:** reglas `templateAccessibility` de angular-eslint en `ng lint`.",
        "**axe:** `@axe-core/playwright` con los tags `wcag2a`, `wcag2aa`, `wcag21a`, `wcag21aa` y `wcag22aa`.",
        "**teclado:** recorrido solo con Tab, Mayús+Tab, Enter, Espacio, flechas y Escape.",
        "**NVDA:** lector de pantalla NVDA con Chrome o Firefox.",
        "**Lighthouse:** categoría Accesibilidad.",
        "**revisión visual:** inspección manual con zoom, ancho de 320 px, espaciado de texto o emulación.",
    ])
    criterios = cargar_wcag()
    faltan = [c["id"] for c in criterios if c["id"] not in CHECKLIST]
    sobran = [k for k in CHECKLIST if k not in {c["id"] for c in criterios}]
    if faltan or sobran:
        raise SystemExit(f"Checklist desalineado con wcag22_a_aa.json: faltan {faltan}, sobran {sobran}")
    filas = []
    for c in criterios:
        aplica, metodo, pantalla = CHECKLIST[c["id"]]
        for m in metodo.split(", "):
            if m not in METODOS:
                raise SystemExit(f"{c['id']}: método no permitido {m!r}")
        filas.append([c["id"], c["nombre"], c["nivel"], aplica, metodo, pantalla])
    d.tabla(["ID", "Criterio", "Nivel", "Aplica", "Método", "Pantalla"], filas, [1.3, 4.4, 1.2, 3.2, 3.0, 4.6])


# ================================================================ 5. i18n

def seccion_i18n(d):
    d.h(1, SECCIONES[4])
    d.p("Alcance decidido en Q8: en Fase 1, la interfaz en español e inglés con cambio de idioma en caliente; en "
        "Fase 2, correos, PDF y asistente M10. El inglés es internacional con ortografía estadounidense. La "
        "implementación es la Fase 5b del plan.")

    d.h(2, "5.1 Configuración")
    d.lista([
        "Librería `@jsverse/transloco` (MIT) con `provideTransloco`: `availableLangs: ['es','en']`, "
        "`defaultLang: 'es'`, `fallbackLang: 'es'` y `reRenderOnLangChange: true`.",
        "Archivos de traducción en `frontend/public/i18n/es.json` y `frontend/public/i18n/en.json`, servidos como "
        "`/i18n/{lang}.json` por `TranslocoHttpLoader`.",
        "`LocaleService` guarda el idioma en `hl_lang`, calcula el locale (`es-CL`, `es-BO` o `en`) y el huso del "
        "país, y actualiza `document.documentElement.lang` sin recargar la página.",
        "El selector ES/EN está en la barra de navegación.",
    ])

    d.h(2, "5.2 Claves")
    d.lista([
        "Formato `modulo.componente.elemento`, en minúsculas y camelCase dentro de cada segmento: "
        "`bl.list.caption`, `payments.form.submit`, `auth.login.password`.",
        "Textos compartidos bajo `common.*` (por ejemplo `common.actions.retry`); los estados de negocio, bajo "
        "`status.<code>`.",
        "Una clave por contexto: no se reutiliza una clave en pantallas distintas aunque el texto coincida.",
        "También son claves los atributos `aria-label`, `alt`, `title` y `placeholder`, y los mensajes en TypeScript "
        "(`translate()`).",
        "Los términos salen del glosario de la sección 6.",
    ])

    d.h(2, "5.3 Formatos por idioma y país (Q8)")
    d.tabla(["Idioma / país", "Locale", "Fecha", "Hora", "Huso horario", "Número", "Montos"], [
        ["Español, Chile", "`es-CL`", "dd-MM-yyyy (05-10-2026)", "24 h (14:30)", "America/Santiago", "1.234.567,89",
         "CLP 1.234.567 (0 decimales); USD 1.234,56"],
        ["Español, Bolivia", "`es-BO`", "dd/MM/yyyy (05/10/2026)", "24 h (14:30)", "America/La_Paz", "1.234.567,89",
         "BOB 1.234,56; USD 1.234,56"],
        ["Inglés", "`en`", "dd MMM yyyy (05 Oct 2026)", "24 h (14:30)", "El del país del usuario",
         "1,234,567.89", "CLP 1,234,567; BOB 1,234.56; EUR 1,234.56"],
    ], [2.6, 1.6, 3.2, 1.8, 2.6, 2.2, 4.0])
    d.lista([
        "Los pipes `hlDate`, `hlNumber` y `hlCurrency` (Fase 5b) usan `Intl.*`, que admite husos IANA, y guardan en "
        "caché los formateadores por locale y opciones.",
        "Las fechas muestran el huso (`timeZoneName: 'short'`). Los plazos se calculan en UTC con calendario de "
        "negocio y se presentan en el huso del país (DC3).",
        "Los montos siempre llevan el código ISO (`currencyDisplay: 'code'`): CLP sin decimales; BOB, USD y EUR con 2.",
        "Los ejemplos de la tabla son ilustrativos; el formato final lo produce `Intl` con el locale indicado.",
    ])

    d.h(2, "5.4 Redacción sin concatenar")
    d.p("Una frase es una sola clave con parámetros. No se arma texto uniendo claves o literales, porque el orden de "
        "las palabras cambia entre idiomas.")
    d.codigo('// Incorrecto\n'
             '"payments.summary.prefix": "Total a pagar:",\n'
             '"payments.summary.suffix": "en"\n'
             '\n'
             '// Correcto\n'
             '"payments.summary.total": "Total a pagar: {{ amount }} en {{ count }} BL"')

    d.h(2, "5.5 Plurales con ICU")
    d.p("Los plurales usan la sintaxis ICU MessageFormat. Transloco la interpreta con el complemento "
        "`@jsverse/transloco-messageformat` (MIT); la Fase 5b debe agregarlo junto con Transloco y registrarlo en "
        "`licencias-dependencias.md`.")
    d.codigo('"bl.list.count": "{count, plural, =0 {Sin BL} one {# BL} other {# BL}}"\n'
             '"notifications.inbox.unread": "{count, plural, =0 {No unread notifications} one {# unread notification} '
             'other {# unread notifications}}"')

    d.h(2, "5.6 Paridad de claves en CI")
    d.lista([
        "`npm run check:i18n` (`frontend/scripts/check-i18n.mjs`): mismas claves en `es.json` y `en.json`, sin "
        "valores vacíos y sin claves usadas que no existan.",
        "`npm run check:i18n-text` (`frontend/scripts/check-hardcoded-text.mjs`): ningún texto ni atributo accesible "
        "literal en las plantillas; las excepciones se declaran en el mismo script.",
        "`npx transloco-keys-manager extract` agrega las claves nuevas al extraer textos.",
        "Ambos chequeos corren en el job `frontend` de `.github/workflows/pr-tests.yml`.",
        "Fuera de alcance por ahora: la traducción de ProblemDetails, correos y PDF del backend (tarea de Fase 2).",
    ])


# ================================================================ 6. glosario

def leer_glosario():
    """(reglas, cabeceras, filas) de docs/requerimientos/glosario-es-en.md."""
    with open(GLOSARIO_MD, encoding="utf-8") as f:
        lineas = f.read().splitlines()
    reglas = [l[2:].strip() for l in lineas if l.startswith("- ")]
    tabla = [l for l in lineas if l.startswith("|")]
    celdas = [[c.strip() for c in l.strip().strip("|").split("|")] for l in tabla]
    cab = celdas[0]
    filas = [c for c in celdas[1:] if not all(re.fullmatch(r"-+", x) for x in c)]
    return reglas, cab, filas


def seccion_glosario(d):
    d.h(1, SECCIONES[5])
    reglas, cab, filas = leer_glosario()
    d.p("Fuente única: `docs/requerimientos/glosario-es-en.md` (Q8). Las claves Transloco y los textos en inglés "
        f"usan estos términos. {len(filas)} términos.")
    d.lista(reglas)
    d.tabla(cab, filas, [5.0, 5.0, 7.0])


# ================================================================ principal

def construir():
    d = Documento(TITULO, "Tokens, tipografía, componentes accesibles, checklist WCAG 2.2 AA e idiomas ES/EN")
    d.p(f"Versión 4.0 · {datetime.date.today().strftime('%d-%m-%Y')} · Preparado para: Hapag-Lloyd Chile y Bolivia · "
        f"Elaborado por: {config.EJECUTA}. Fuente: Fase 3 del plan de actualización de requerimientos Portal 2.0; "
        "decisiones Q8, Q9, DC2, DC3 y DC4 del registro de decisiones v4. Generado por "
        "`scripts/requerimientos/generar_guia.py`; no editar a mano.")
    seccion_tokens(d)
    seccion_tipografia(d)
    seccion_componentes(d)
    seccion_checklist(d)
    seccion_i18n(d)
    seccion_glosario(d)
    return d


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    d = construir()
    escribir_md(SALIDA_MD, a_markdown(d))
    salida = a_docx(d, config.ruta_v4(NOMBRE_DOCX), pie="Guía UI, accesibilidad e idiomas · Portal 2.0 v4")
    print(f"{os.path.relpath(SALIDA_MD, REPO)}")
    print(salida)
    print(f"{len(TOKENS)} tokens + --hl-focus; {len(PARES)} pares de contraste; {len(cargar_wcag())} criterios WCAG")
    return 0


if __name__ == "__main__":
    sys.exit(main())
