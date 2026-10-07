"""Exporta docs/requerimientos/registro-decisiones-v4.md a katu\\v4\\Registro_Decisiones_v4.xlsx (una hoja por tabla)."""
import os
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402
from linea_base import DOCS  # noqa: E402

HOJAS = ["Decisiones", "Persona-Area", "Complementarias", "Calendario", "Publicados"]


def tablas_md(texto):
    tablas, actual = [], []
    for linea in texto.splitlines():
        if linea.startswith("|"):
            celdas = [c.strip() for c in linea.strip().strip("|").split("|")]
            if all(re.fullmatch(r"-+", c) for c in celdas):
                continue
            actual.append([re.sub(r"\*\*|`", "", c) for c in celdas])
        elif actual:
            tablas.append(actual)
            actual = []
    if actual:
        tablas.append(actual)
    return tablas


def main():
    with open(os.path.join(DOCS, "registro-decisiones-v4.md"), encoding="utf-8") as f:
        tablas = tablas_md(f.read())
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, filas in zip(HOJAS, tablas):
        ws = wb.create_sheet(nombre)
        for fila in filas:
            ws.append(fila)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0B2C5C")
        for row in ws.iter_rows():
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, col in enumerate(ws.columns, 1):
            largo = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(max(12, largo * 0.9), 90)
        ws.freeze_panes = "A2"
    salida = config.ruta_v4("Registro_Decisiones_v4.xlsx")
    wb.save(salida)
    print(f"{len(tablas)} tablas -> {salida}")


if __name__ == "__main__":
    main()
