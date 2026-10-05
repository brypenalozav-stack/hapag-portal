import sys,os,io,hashlib
from docx import Document
import openpyxl
K=r"C:\Users\klaze\Desktop\katu"; out=sys.argv[1]
for f in sorted(os.listdir(K)):
    p=os.path.join(K,f); print(f, os.path.getsize(p), hashlib.md5(open(p,'rb').read()).hexdigest()[:8])
    o=io.open(os.path.join(out,f+'.txt'),'w',encoding='utf-8')
    if f.endswith('.docx'):
        d=Document(p)
        body=d.element.body
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        for el in body.iterchildren():
            if el.tag.endswith('}p'):
                para=Paragraph(el,d); t=para.text.strip()
                if t: o.write(f"[{para.style.name}] {t}\n")
            elif el.tag.endswith('}tbl'):
                tb=Table(el,d); o.write("<TABLE>\n")
                for r in tb.rows:
                    cells=[];prev=None
                    for c in r.cells:
                        if c._tc is prev: continue
                        prev=c._tc; cells.append(c.text.strip().replace('\n',' / '))
                    o.write(' | '.join(cells)+'\n')
                o.write("</TABLE>\n")
    else:
        wb=openpyxl.load_workbook(p,data_only=True)
        for ws in wb.worksheets:
            o.write(f"### SHEET {ws.title} ({ws.max_row}x{ws.max_column})\n")
            for i,row in enumerate(ws.iter_rows(values_only=True),1):
                vals=['' if v is None else str(v).replace('\n',' / ') for v in row]
                while vals and not vals[-1]: vals.pop()
                if any(vals): o.write(f"{i}: "+' | '.join(vals)+'\n')
    o.close()
