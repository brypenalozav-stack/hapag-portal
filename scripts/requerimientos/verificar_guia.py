"""Fase 3: verifica la Guía UI/a11y/i18n (Markdown y Word). Sale con 1 si algo falla.

- Las 6 secciones existen, en orden, en el .md y en el .docx.
- El checklist cubre exactamente los ids de wcag22_a_aa.json (55) y ningún método está vacío ni fuera de
  eslint / axe / teclado / NVDA / Lighthouse / revisión visual.
- Los tokens coinciden con frontend/src/styles/_tokens.scss si existe; si no, con la tabla de la Fase 5a §2
  del plan (copiada abajo en TOKENS_FASE_5A). Cada contraste publicado coincide con el recalculado.
- El .md y el .docx tienen los mismos títulos.
- 0 términos prohibidos (terminos.py) en ambos archivos.
"""
import json
import os
import re
import sys

import docx
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402
import terminos  # noqa: E402
from linea_base import DOCS, REPO  # noqa: E402

MD = os.path.join(DOCS, "guia-ui-a11y-i18n.md")
DOCX = os.path.join(config.V4, "Guia_UI_Accesibilidad_i18n_v4.docx")
WCAG_JSON = os.path.join(os.path.dirname(os.path.abspath(__file__)), "wcag22_a_aa.json")
TOKENS_SCSS = os.path.join(REPO, "frontend", "src", "styles", "_tokens.scss")

SECCIONES = [
    "1. Tokens de diseño",
    "2. Tipografía",
    "3. Componentes y patrones",
    "4. Checklist WCAG 2.2 AA",
    "5. Estrategia i18n con Transloco",
    "6. Glosario ES/EN",
]
METODOS = {"eslint", "axe", "teclado", "NVDA", "Lighthouse", "revisión visual"}

# Fase 5a §2 del plan: lista cerrada de tokens (DC4).
TOKENS_FASE_5A = {
    "$hl-dark": "#33424f",
    "$hl-orange": "#ff6600",
    "$hl-orange-700": "#b84a00",
    "$hl-green": "#009840",
    "$hl-green-700": "#007a33",
    "$hl-blue": "#004d6c",
    "$hl-light-gray": "#f5f7fa",
    "$hl-border": "#e2e8f0",
    "$hl-body-color": "#2d3748",
    "$hl-white": "#ffffff",
    "$hl-dark-bg": "#1a202c",
    "$hl-dark-text": "#e2e8f0",
    "$hl-surface-muted": "#f8fafc",
    "$hl-text-muted": "#4a5568",
}
FOCO_FASE_5A = ("--hl-focus", "#004d6c")  # --hl-focus: $hl-blue

RE_HEX = re.compile(r"#[0-9a-fA-F]{6}\b")


def tokens_esperados():
    if os.path.exists(TOKENS_SCSS):
        with open(TOKENS_SCSS, encoding="utf-8") as f:
            pares = re.findall(r"^\s*(\$hl-[\w-]+)\s*:\s*(#[0-9a-fA-F]{3,8})", f.read(), re.M)
        return {k: v.lower() for k, v in pares}, os.path.relpath(TOKENS_SCSS, REPO)
    return dict(TOKENS_FASE_5A), "tabla de la Fase 5a §2 del plan"


def _lineal(c):
    c = c / 255
    return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4


def contraste(a, b):
    def lum(h):
        h = h.lstrip("#")
        r, g, bb = (int(h[i:i + 2], 16) for i in (0, 2, 4))
        return 0.2126 * _lineal(r) + 0.7152 * _lineal(g) + 0.0722 * _lineal(bb)
    x, y = sorted((lum(a), lum(b)), reverse=True)
    return (x + 0.05) / (y + 0.05)


def limpiar(celda):
    return re.sub(r"`|\*\*", "", celda).strip()


# ---------------------------------------------------------------- lectura

def leer_md():
    with open(MD, encoding="utf-8") as f:
        texto = f.read()
    titulos = [re.sub(r"^#+\s*", "", l).strip() for l in texto.splitlines() if re.match(r"^#{2,3}\s", l)]
    tablas, actual = [], []
    for l in texto.splitlines():
        if l.startswith("|"):
            celdas = [c.strip() for c in re.split(r"(?<!\\)\|", l.strip())[1:-1]]
            if not all(re.fullmatch(r"-+", c) for c in celdas):
                actual.append([limpiar(c.replace("\\|", "|")) for c in celdas])
        elif actual:
            tablas.append(actual)
            actual = []
    if actual:
        tablas.append(actual)
    return texto, titulos, tablas


def leer_docx():
    d = docx.Document(DOCX)
    titulos = [p.text.strip() for p in d.paragraphs if p.style.name in ("Heading 1", "Heading 2")]
    tablas = []
    for t in d.tables:
        tablas.append([[c.text.strip() for c in fila.cells] for fila in t.rows])
    partes = ["".join(x.text or "" for x in p.iter(qn("w:t"))) for p in d.element.body.iter(qn("w:p"))]
    for s in d.sections:
        partes += [p.text for p in s.footer.paragraphs]
    return "\n".join(partes), titulos, tablas


def tabla_por_cabecera(tablas, primera):
    for t in tablas:
        if t and t[0] and t[0][0] == primera:
            return t
    return None


# ---------------------------------------------------------------- comprobaciones

def verificar_secciones(nombre, titulos, errores):
    nivel1 = [t for t in titulos if re.match(r"^\d+\.\s", t)]
    if nivel1 != SECCIONES:
        errores.append(f"{nombre}: secciones {nivel1} (se esperaban {SECCIONES})")


def verificar_checklist(nombre, tablas, ids, errores):
    t = tabla_por_cabecera(tablas, "ID")
    if t is None:
        errores.append(f"{nombre}: no hay tabla de checklist (cabecera 'ID')")
        return 0
    cab = t[0]
    for requerida in ("ID", "Criterio", "Aplica", "Método", "Pantalla"):
        if requerida not in cab:
            errores.append(f"{nombre}: el checklist no tiene la columna {requerida!r}")
            return 0
    i_id, i_met, i_pan = cab.index("ID"), cab.index("Método"), cab.index("Pantalla")
    vistos = [f[i_id] for f in t[1:]]
    if sorted(vistos) != sorted(ids) or len(vistos) != len(set(vistos)):
        faltan = sorted(set(ids) - set(vistos))
        sobran = sorted(set(vistos) - set(ids))
        errores.append(f"{nombre}: el checklist no cubre exactamente los {len(ids)} criterios "
                       f"(faltan {faltan}, sobran {sobran}, filas {len(vistos)})")
    for f in t[1:]:
        metodos = [m.strip() for m in f[i_met].split(",") if m.strip()]
        if not metodos:
            errores.append(f"{nombre}: {f[i_id]} sin método")
        for m in metodos:
            if m not in METODOS:
                errores.append(f"{nombre}: {f[i_id]} con método no permitido {m!r}")
        if not f[i_pan].strip():
            errores.append(f"{nombre}: {f[i_id]} sin pantalla")
    return len(vistos)


def verificar_tokens(nombre, tablas, esperados, errores):
    t = tabla_por_cabecera(tablas, "Token Sass")
    if t is None:
        errores.append(f"{nombre}: no hay tabla de tokens (cabecera 'Token Sass')")
        return
    cab = t[0]
    i_tok, i_css, i_val = cab.index("Token Sass"), cab.index("Variable CSS"), cab.index("Valor")
    publicados = {}
    foco = None
    for f in t[1:]:
        if f[i_tok].startswith("$"):
            publicados[f[i_tok]] = f[i_val].lower()
            if f[i_css] != "--" + f[i_tok][1:]:
                errores.append(f"{nombre}: {f[i_tok]} con variable CSS {f[i_css]!r}")
        elif f[i_css] == FOCO_FASE_5A[0]:
            foco = f[i_val].lower()
    if publicados != esperados:
        faltan = sorted(set(esperados) - set(publicados))
        sobran = sorted(set(publicados) - set(esperados))
        distintos = sorted(k for k in set(esperados) & set(publicados) if esperados[k] != publicados[k])
        errores.append(f"{nombre}: tokens distintos de la fuente (faltan {faltan}, sobran {sobran}, "
                       f"valor distinto {distintos})")
    esperado_foco = esperados.get("$hl-blue", FOCO_FASE_5A[1])
    if foco != esperado_foco:
        errores.append(f"{nombre}: --hl-focus es {foco!r} (se esperaba {esperado_foco})")


def verificar_contrastes(nombre, tablas, errores):
    t = tabla_por_cabecera(tablas, "Tema")
    if t is None:
        errores.append(f"{nombre}: no hay tabla de contrastes (cabecera 'Tema')")
        return 0
    cab = t[0]
    i_fg, i_bg, i_cr = cab.index("Primer plano"), cab.index("Fondo"), cab.index("Contraste")
    i_aa, i_3 = cab.index("Texto normal (4,5:1)"), cab.index("Texto grande y no textual (3:1)")
    for f in t[1:]:
        fg, bg = RE_HEX.search(f[i_fg]), RE_HEX.search(f[i_bg])
        if not fg or not bg:
            errores.append(f"{nombre}: par de contraste sin color: {f}")
            continue
        x = contraste(fg.group(0), bg.group(0))
        if f[i_cr] != f"{x:.2f}".replace(".", ",") + ":1":
            errores.append(f"{nombre}: contraste {fg.group(0)} sobre {bg.group(0)} publicado {f[i_cr]}, "
                           f"calculado {x:.2f}")
        if (f[i_aa] == "Cumple") != (x >= 4.5) or (f[i_3] == "Cumple") != (x >= 3):
            errores.append(f"{nombre}: resultado AA incorrecto para {fg.group(0)} sobre {bg.group(0)}")
    return len(t) - 1


def verificar_terminos(nombre, texto, errores):
    hits = terminos.encontrar_prohibidos(texto)
    if hits:
        errores.append(f"{nombre}: {len(hits)} término(s) prohibido(s)")


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    errores = []
    for p in (MD, DOCX):
        if not os.path.exists(p):
            errores.append(f"no existe {p}")
    if errores:
        for e in errores:
            print("ERROR", e)
        return 1
    with open(WCAG_JSON, encoding="utf-8") as f:
        criterios = json.load(f)["criterios"]
    ids = [c["id"] for c in criterios]
    niveles = [c["nivel"] for c in criterios]
    if len(ids) != 55 or len(set(ids)) != 55 or niveles.count("A") != 31 or niveles.count("AA") != 24 \
            or "4.1.1" in ids:
        errores.append(f"wcag22_a_aa.json: {len(ids)} criterios ({niveles.count('A')} A, {niveles.count('AA')} AA); "
                       "se esperaban 55 (31 A + 24 AA) sin 4.1.1")
    esperados, fuente = tokens_esperados()

    texto_md, titulos_md, tablas_md = leer_md()
    texto_dx, titulos_dx, tablas_dx = leer_docx()
    resumen = {}
    for nombre, titulos, tablas, texto in (("md", titulos_md, tablas_md, texto_md),
                                           ("docx", titulos_dx, tablas_dx, texto_dx)):
        verificar_secciones(nombre, titulos, errores)
        n = verificar_checklist(nombre, tablas, ids, errores)
        verificar_tokens(nombre, tablas, esperados, errores)
        pares = verificar_contrastes(nombre, tablas, errores)
        verificar_terminos(nombre, texto, errores)
        resumen[nombre] = (len(titulos), n, pares)
    if titulos_md != titulos_dx:
        solo_md = [t for t in titulos_md if t not in titulos_dx]
        solo_dx = [t for t in titulos_dx if t not in titulos_md]
        errores.append(f"títulos distintos entre md y docx (solo md: {solo_md}; solo docx: {solo_dx})")

    print(f"tokens comparados con: {fuente} ({len(esperados)} tokens + --hl-focus)")
    for nombre, (nt, n, pares) in resumen.items():
        print(f"{nombre}: {nt} títulos, {n} criterios en el checklist, {pares} pares de contraste")
    for e in errores:
        print("ERROR", e)
    if errores:
        print(f"{len(errores)} problema(s)")
        return 1
    print("guía correcta")
    return 0


if __name__ == "__main__":
    sys.exit(main())
