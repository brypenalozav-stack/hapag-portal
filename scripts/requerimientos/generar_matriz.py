"""Fase 3: Matriz de trazabilidad Portal 2.0 v4 (una fila por ficha de la especificación v4).

Fuentes (no usa el Gantt):
- katu\\v4\\Portal_2.0_Especificacion_Funcional_v4.docx: las 134 fichas (docx_utils.fichas).
- katu\\v4\\Portal_2_0_Pendientes_v4.xlsx: tareas por las columnas "Ficha spec v4" y "Contrato de integración",
  y por mención del ID de la ficha en la tarea o el comentario.
- katu\\v4\\Funcionalidades NexusV2_v4.xlsx: función NexusV2 por la columna "Ficha spec v4".
- docs/integraciones: fichas de cada contrato (inventario del README, campo "Fichas" del contrato y la asignación
  de la Fase 2), `info.x-estado` / fila "Estado" y `info.x-puerto` / fila "Puerto".
- backend/src: la interfaz del puerto y las clases Dummy*/Http* que la implementan (si la Fase 2 asignó puertos
  concretos a la ficha, solo esos).
- katu\\v4\\Matriz_Fichas_Linea_Base_v4.xlsx: cobertura C/P/N de la línea base (las fichas nuevas son N).

Salidas: docs/requerimientos/matriz-trazabilidad.md y katu\\v4\\Matriz_Trazabilidad_v4.docx (horizontal, con la
fila de cabecera repetida en cada página).
"""
import datetime
import glob
import hashlib
import os
import re
import subprocess
import sys

import docx
import openpyxl
import yaml

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402
import docx_utils  # noqa: E402
import editar_operativos_v4 as ed  # noqa: E402
from generar_guia import Documento, a_docx, a_markdown, escribir_md  # noqa: E402
from linea_base import DOCS, REPO  # noqa: E402

ESPEC_V4 = os.path.join(config.V4, "Portal_2.0_Especificacion_Funcional_v4.docx")
PENDIENTES_V4 = os.path.join(config.V4, "Portal_2_0_Pendientes_v4.xlsx")
NEXUS_V4 = os.path.join(config.V4, "Funcionalidades NexusV2_v4.xlsx")
MATRIZ_BASE = os.path.join(config.V4, "Matriz_Fichas_Linea_Base_v4.xlsx")
INTEGRACIONES = os.path.join(REPO, "docs", "integraciones")
README = os.path.join(INTEGRACIONES, "README.md")
BACKEND_SRC = os.path.join(REPO, "backend", "src")
SALIDA_MD = os.path.join(DOCS, "matriz-trazabilidad.md")
NOMBRE_DOCX = "Matriz_Trazabilidad_v4.docx"

TITULO = "Matriz de trazabilidad — Portal 2.0 v4"
CABECERAS = ["Ficha", "Fase", "Tareas Pendientes_v4", "Función NexusV2", "Contrato", "Estado contrato",
             "Código hapag-portal", "Cobertura"]
ANCHOS = [4.2, 1.0, 3.4, 2.6, 2.0, 4.6, 7.2, 1.7]
SIN_TAREA = "Sin tarea"
SIN_CONTRATO = "Sin contrato"

RE_FICHA = re.compile(r"(?<![\w-])((?:M\d+|NF)-\d{2})(?!\d)")
RE_RANGO = re.compile(r"(?<![\w-])((?:M\d+|NF))-(\d{2}) a \1-(\d{2})")
RE_CONTRATO = re.compile(r"\bCT-[A-Z]+\b")


def ids_fichas(texto):
    """IDs de ficha citados en un texto, con los rangos 'M10-01 a M10-06' expandidos."""
    texto = str(texto or "")
    ids = set(RE_FICHA.findall(texto))
    for pref, a, b in RE_RANGO.findall(texto):
        ids.update(f"{pref}-{n:02d}" for n in range(int(a), int(b) + 1))
    return ids


# ---------------------------------------------------------------- Pendientes y NexusV2

def tareas_pendientes():
    """[(id, tarea, comentario, fichas, contratos)] de Pendientes_v4."""
    ws = openpyxl.load_workbook(PENDIENTES_V4, read_only=True)[ed.HOJA_PENDIENTES]
    filas = list(ws.iter_rows(min_row=4, values_only=True))
    cab = list(filas[0])
    i_tarea, i_com = cab.index("Tarea"), cab.index("Comentarios")
    i_ficha, i_ct = cab.index("Ficha spec v4"), cab.index("Contrato de integración")
    out = []
    for r in filas[1:]:
        if r[0] is None:
            continue
        out.append((int(r[0]), str(r[i_tarea] or ""), str(r[i_com] or ""), ids_fichas(r[i_ficha]),
                    set(RE_CONTRATO.findall(str(r[i_ct] or "")))))
    return out


def funciones_nexus():
    """{ficha: [funciones NexusV2]} desde la columna 'Ficha spec v4' de NexusV2_v4."""
    ws = openpyxl.load_workbook(NEXUS_V4, read_only=True).active
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    cab = list(filas[0])
    i_ficha = cab.index("Ficha spec v4")
    out = {}
    for r in filas[1:]:
        if not r[0]:
            continue
        nombre = str(r[0]).split(" (")[0].strip()
        for fid in ids_fichas(r[i_ficha]):
            out.setdefault(fid, []).append(nombre)
    return out


# ---------------------------------------------------------------- contratos (docs/integraciones)

def fila_md(texto, campo):
    m = re.search(rf"^\|\s*{re.escape(campo)}\s*\|\s*(.+?)\s*\|\s*$", texto, re.M)
    return m.group(1).strip() if m else ""


def contratos():
    """{CT-ID: {archivo, estado, puertos[(interfaz, clave)], fichas}} leídos de los archivos de contrato."""
    out = {}
    for cid, archivo, _ in ed.CONTRATOS:
        ruta = os.path.join(INTEGRACIONES, archivo)
        with open(ruta, encoding="utf-8") as f:
            texto = f.read()
        if archivo.endswith(".yaml"):
            info = yaml.safe_load(texto)["info"]
            if info.get("x-contrato") not in (None, cid):
                raise SystemExit(f"{archivo}: x-contrato {info.get('x-contrato')} distinto de {cid}")
            estado = str(info.get("x-estado") or "").strip()
            puertos = []
            for p in info.get("x-puerto") or []:
                m = re.match(r'^(I\w+)(?:\("(\w+)"\))?$', str(p).strip())
                if not m:
                    raise SystemExit(f"{archivo}: x-puerto no reconocido {p!r}")
                puertos.append((m.group(1), m.group(2)))
            m = re.search(r"Fichas?:([^\n]*)", str(info.get("description") or ""))
            fichas = ids_fichas(m.group(1)) if m else set()
        else:
            estado = fila_md(texto, "Estado")
            fila_puerto = fila_md(texto, "Puerto")
            puertos = ([] if fila_puerto.startswith("Sin puerto")
                       else [(i, None) for i in re.findall(r"`(I\w+)`", fila_puerto)])
            fichas = ids_fichas(fila_md(texto, "Fichas"))
        if not estado:
            raise SystemExit(f"{archivo}: sin estado")
        out[cid] = {"archivo": archivo, "estado": estado, "puertos": puertos, "fichas": fichas}
    return out


def fichas_readme():
    """{CT-ID: fichas} del inventario de docs/integraciones/README.md (columnas Fichas y Estado)."""
    with open(README, encoding="utf-8") as f:
        lineas = [l for l in f.read().splitlines() if l.startswith("|")]
    cab = [c.strip() for c in lineas[0].strip().strip("|").split("|")]
    i_fichas, i_estado = cab.index("Fichas"), cab.index("Estado")
    out = {}
    for l in lineas[2:]:
        celdas = [c.strip() for c in l.strip().strip("|").split("|")]
        if len(celdas) != len(cab):
            break
        for cid in RE_CONTRATO.findall(celdas[i_estado]):
            out.setdefault(cid, set()).update(ids_fichas(celdas[i_fichas]))
    return out


def contratos_por_ficha(cts):
    out = {}
    for cid, fichas in fichas_readme().items():
        for fid in fichas:
            out.setdefault(fid, set()).add(cid)
    for cid, c in cts.items():
        for fid in c["fichas"]:
            out.setdefault(fid, set()).add(cid)
    return out


def contratos_fase2(fid):
    """Contratos asignados a la ficha en la Fase 2 (editar_operativos_v4.contrato_ficha)."""
    return set(RE_CONTRATO.findall(ed.contrato_ficha(fid)))


# ---------------------------------------------------------------- código (backend/src)

RE_INTERFAZ = re.compile(r"\binterface\s+(I\w+)")
RE_CLASE = re.compile(r"\bclass\s+((?:Dummy|Http)\w+)\s*(?:\([^)]*\))?\s*:\s*([^{]+)\{", re.S)


def indice_codigo():
    """(interfaces {nombre: ruta}, clases [(nombre, ruta, {interfaces})]) con rutas relativas a backend/src."""
    interfaces, clases = {}, []
    for ruta in glob.glob(os.path.join(BACKEND_SRC, "**", "*.cs"), recursive=True):
        if os.sep + "obj" + os.sep in ruta or os.sep + "bin" + os.sep in ruta:
            continue
        with open(ruta, encoding="utf-8-sig") as f:
            texto = f.read()
        rel = os.path.relpath(ruta, BACKEND_SRC).replace(os.sep, "/")
        for n in RE_INTERFAZ.findall(texto):
            interfaces.setdefault(n, rel)
        for nombre, bases in RE_CLASE.findall(texto):
            impl = {re.sub(r"<.*", "", b).strip() for b in bases.split(",")}
            clases.append((nombre, rel, impl))
    return interfaces, clases


def puertos_de_ficha(fid, cid, contrato):
    """Puertos del contrato que usa la ficha. Si la Fase 2 asignó puertos a la ficha (PUERTOS_FICHA), solo esos;
    si no, todos los del contrato."""
    usados = {ed.PUERTO[p][0] for p in ed.PUERTOS_FICHA.get(fid, []) if cid in ed.PUERTO[p][3].split(", ")}
    propios = [p for p in contrato["puertos"] if p[0] in usados]
    return propios or contrato["puertos"]


def codigo_contrato(puertos, interfaces, clases):
    """Rutas (relativas a backend/src) de la interfaz y de las clases Dummy* y Http* que implementan los puertos."""
    rutas = []
    for interfaz, clave in puertos:
        if interfaz not in interfaces:
            raise SystemExit(f"Puerto {interfaz} sin interfaz en backend/src")
        rutas.append(interfaces[interfaz])
        for prefijo in ("Dummy", "Http"):
            for nombre, ruta, impl in sorted(clases):
                if not nombre.startswith(prefijo) or interfaz not in impl:
                    continue
                if prefijo == "Http" and clave and clave.lower() not in nombre.lower():
                    continue
                rutas.append(ruta)
    return list(dict.fromkeys(rutas))


# ---------------------------------------------------------------- filas

def construir_filas():
    d = docx.Document(ESPEC_V4)
    fichas = docx_utils.fichas(d)
    if len(fichas) != 134:
        raise SystemExit(f"La especificación v4 tiene {len(fichas)} fichas (se esperaban 134)")
    tareas = tareas_pendientes()
    funciones = funciones_nexus()
    cts = contratos()
    por_ficha = contratos_por_ficha(cts)
    interfaces, clases = indice_codigo()
    base = ed.matriz_base()

    validar = {}
    for tid, texto, _, _, cs in tareas:
        if texto.startswith("Validar contrato"):
            for cid in cs:
                validar.setdefault(cid, []).append(tid)

    filas = []
    for fid, titulo, fase in fichas:
        directas = [tid for tid, texto, com, fs, _ in tareas
                    if fid in fs or fid in ids_fichas(texto) or fid in ids_fichas(com)]
        ids_ct = sorted(por_ficha.get(fid, set()) | contratos_fase2(fid),
                        key=[c[0] for c in ed.CONTRATOS].index)
        celda_tareas = []
        if directas:
            celda_tareas.append(", ".join(str(t) for t in sorted(directas)))
        via = [f"{t} ({cid})" for cid in ids_ct for t in validar.get(cid, [])]
        if via:
            celda_tareas.append("Validación de contrato: " + ", ".join(via))
        estado = [f"{cts[c]['estado']} ({c})" for c in ids_ct]
        codigo = []
        for c in ids_ct:
            codigo += codigo_contrato(puertos_de_ficha(fid, c, cts[c]), interfaces, clases)
        codigo = list(dict.fromkeys(codigo))
        cobertura = base[fid][0] if fid in base else "N"
        filas.append({
            "id": fid, "fase": ed.fase_de(fase),
            "celdas": [f"{fid} {titulo}", ed.fase_de(fase), celda_tareas or [SIN_TAREA],
                       ", ".join(dict.fromkeys(funciones.get(fid, []))) or "—",
                       ", ".join(ids_ct) or SIN_CONTRATO, estado or [SIN_CONTRATO], codigo or ["—"], cobertura],
            "nueva": fid not in base,
        })
    return filas, cts


# ---------------------------------------------------------------- documento

def git(*args):
    return subprocess.check_output(["git", "-C", REPO, *args], text=True).strip()


def sha12(ruta):
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(1 << 20), b""):
            h.update(bloque)
    return h.hexdigest()[:12]


def construir():
    filas, cts = construir_filas()
    ahora = datetime.datetime.now()
    commit = git("rev-parse", "HEAD")
    rama = git("rev-parse", "--abbrev-ref", "HEAD")
    cambios = git("status", "--porcelain", "--", "backend/src", "docs/integraciones")

    d = Documento(TITULO, "Fichas de la especificación v4 → tareas, NexusV2, contratos de integración y código")
    d.p(f"Versión 4.0 · Preparado para: Hapag-Lloyd Chile y Bolivia · Elaborado por: {config.EJECUTA}. "
        "Generado por `scripts/requerimientos/generar_matriz.py`; no editar a mano.")

    d.h(1, "1. Fecha y commit de generación")
    d.lista([
        f"Fecha: {ahora.strftime('%d-%m-%Y %H:%M')} (hora local del equipo que la generó).",
        f"Commit de `hapag-portal`: `{commit}` (rama `{rama}`)."
        + (" Hay cambios sin commit en `backend/src` o `docs/integraciones`." if cambios else ""),
        f"Especificación: `{os.path.basename(ESPEC_V4)}` (SHA-256 `{sha12(ESPEC_V4)}…`).",
        f"Pendientes: `{os.path.basename(PENDIENTES_V4)}` (SHA-256 `{sha12(PENDIENTES_V4)}…`).",
        f"NexusV2: `{os.path.basename(NEXUS_V4)}` (SHA-256 `{sha12(NEXUS_V4)}…`).",
        f"Cobertura de línea base: `{os.path.basename(MATRIZ_BASE)}` (SHA-256 `{sha12(MATRIZ_BASE)}…`).",
        f"Contratos: `docs/integraciones/` ({len(cts)} contratos). Código: `backend/src`.",
        "El Gantt de macros no es fuente de esta matriz (fuera del alcance).",
    ])

    d.h(1, "2. Criterios de cada columna")
    d.lista([
        "**Ficha:** ID y título del encabezado de la ficha en la especificación v4.",
        "**Fase:** fase según el encabezado de la ficha (Q1: manda el encabezado).",
        "**Tareas Pendientes_v4:** IDs de las tareas cuya columna \"Ficha spec v4\" cita la ficha o cuya tarea o "
        "comentario menciona su ID. Aparte, las tareas \"Validar contrato\" de cada contrato de la ficha. "
        f"\"{SIN_TAREA}\" si no hay ninguna.",
        "**Función NexusV2:** funciones de `Funcionalidades NexusV2_v4.xlsx` cuya columna \"Ficha spec v4\" cita la "
        "ficha.",
        "**Contrato:** IDs de la Fase 6a asociados a la ficha en el inventario `docs/integraciones/README.md`, en el "
        "propio contrato o en la columna \"Contrato de integración\" de la Fase 2.",
        "**Estado contrato:** `info.x-estado` del contrato OpenAPI, o la fila \"Estado\" de los contratos en "
        f"Markdown. \"{SIN_CONTRATO}\" si la ficha no tiene contrato.",
        "**Código hapag-portal:** rutas relativas a `backend/src` de la interfaz del puerto del contrato "
        "(`info.x-puerto` o fila \"Puerto\") y de las clases `Dummy*` y `Http*` que la implementan. Si la Fase 2 "
        "asignó puertos concretos a la ficha, solo esos. \"—\" si el contrato no tiene puerto o la ficha no tiene "
        "contrato.",
        "**Cobertura:** C (cubierta), P (parcial) o N (no existe) según la línea base de la Fase 0; las fichas "
        "nuevas de v4 (M11-xx, M2-10 y M8-09) son N.",
    ])

    d.h(1, "3. Resumen")
    resumen = []
    for fase in ("0", "1", "2"):
        fs = [f for f in filas if f["fase"] == fase]
        cob = [f["celdas"][7] for f in fs]
        resumen.append([f"Fase {fase}", str(len(fs)),
                        str(sum(1 for f in fs if f["celdas"][2] != [SIN_TAREA])),
                        str(sum(1 for f in fs if f["celdas"][4] != SIN_CONTRATO)),
                        str(sum(1 for f in fs if f["celdas"][6] != ["—"])),
                        f"{cob.count('C')} / {cob.count('P')} / {cob.count('N')}"])
    resumen.append(["Total", str(len(filas)), str(sum(1 for f in filas if f["celdas"][2] != [SIN_TAREA])),
                    str(sum(1 for f in filas if f["celdas"][4] != SIN_CONTRATO)),
                    str(sum(1 for f in filas if f["celdas"][6] != ["—"])),
                    " / ".join(str([f["celdas"][7] for f in filas].count(x)) for x in "CPN")])
    d.tabla(["Fase", "Fichas", "Con tarea", "Con contrato", "Con código de integración", "Cobertura C / P / N"],
            resumen, [2.5, 2.0, 2.5, 2.5, 4.0, 4.0])
    sin_tarea_f1 = [f["id"] for f in filas if f["fase"] == "1" and f["celdas"][2] == [SIN_TAREA]]
    if sin_tarea_f1:
        d.p(f"Fichas de Fase 1 sin tarea en Pendientes_v4 ({len(sin_tarea_f1)}): {', '.join(sin_tarea_f1)}.")
    else:
        d.p("Todas las fichas de Fase 1 tienen al menos una tarea en Pendientes_v4.")

    d.h(1, "4. Matriz")
    d.tabla(CABECERAS, [f["celdas"] for f in filas], ANCHOS)
    return d, filas


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    d, filas = construir()
    escribir_md(SALIDA_MD, a_markdown(d))
    salida = a_docx(d, config.ruta_v4(NOMBRE_DOCX), horizontal=True, tam_tabla=8,
                    pie="Matriz de trazabilidad · Portal 2.0 v4")
    print(os.path.relpath(SALIDA_MD, REPO))
    print(salida)
    con_ct = sum(1 for f in filas if f["celdas"][4] != SIN_CONTRATO)
    print(f"{len(filas)} fichas; {con_ct} con contrato; "
          f"{sum(1 for f in filas if f['celdas'][2] == [SIN_TAREA] and f['fase'] == '1')} de Fase 1 sin tarea")
    return 0


if __name__ == "__main__":
    sys.exit(main())
