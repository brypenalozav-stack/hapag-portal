"""Utilidades de edición de la especificación (.docx) con python-docx.

El índice de la especificación es un campo `TOC \\f P \\h \\l "1-2"` que se arma con
campos TC (`TC "<texto>" \\f P \\l <nivel>`), no con los estilos de título. Por eso
toda inserción o renumeración de títulos debe mantener sus campos TC.
El .docx solo se guarda con python-docx; LibreOffice nunca lo reescribe.
"""
import copy
import os
import re
import shutil
import tempfile
import zipfile

from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph

RE_TC = re.compile(r'^\s*TC\s+"(.*)"\s+\\f\s+P\s+\\l\s+(\d+)\s*$', re.S)
RE_FICHA = re.compile(r"^((?:M\d+|NF)-\d+)\s+")


# ---------------------------------------------------------------- recorrido

def iter_bloques(doc):
    """Párrafos y tablas del cuerpo, en orden de documento."""
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            yield Paragraph(el, doc)
        elif el.tag == qn("w:tbl"):
            yield Table(el, doc)


def iter_parrafos(doc):
    for b in iter_bloques(doc):
        if isinstance(b, Paragraph):
            yield b


# ---------------------------------------------------------------- campos TC

def _campos_tc(p_el):
    """Campos TC de un párrafo: lista de (instrText_elements, texto, nivel)."""
    campos = []
    dentro = False
    instr = []
    for el in p_el.iter():
        if el.tag == qn("w:fldChar"):
            tipo = el.get(qn("w:fldCharType"))
            if tipo == "begin":
                dentro, instr = True, []
            elif tipo == "end" and dentro:
                dentro = False
                m = RE_TC.match("".join(i.text or "" for i in instr))
                if m:
                    campos.append((instr, m.group(1), int(m.group(2))))
        elif el.tag == qn("w:instrText") and dentro:
            instr.append(el)
    return campos


def list_tc_fields(doc):
    """[(texto, nivel)] de todos los campos TC del documento, en orden."""
    out = []
    for p in iter_parrafos(doc):
        for _, texto, nivel in _campos_tc(p._p):
            out.append((texto, nivel))
    return out


def update_tc_fields(doc, old, new):
    """Reemplaza el texto del campo TC cuyo texto es exactamente `old`. Devuelve cuántos cambió."""
    n = 0
    for p in iter_parrafos(doc):
        for instr, texto, nivel in _campos_tc(p._p):
            if texto == old:
                instr[0].text = f' TC "{new}" \\f P \\l {nivel} '
                instr[0].set(qn("xml:space"), "preserve")
                for extra in instr[1:]:
                    extra.text = ""
                n += 1
    return n


def _run(hijo):
    r = OxmlElement("w:r")
    r.append(hijo)
    return r


def insert_tc_field(paragraph, text, level):
    """Inserta al inicio del párrafo un campo TC con la misma estructura que los existentes."""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f' TC "{text}" \\f P \\l {level} '
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    p = paragraph._p
    pos = 1 if (len(p) and p[0].tag == qn("w:pPr")) else 0
    for i, el in enumerate([_run(begin), _run(instr), _run(end)]):
        p.insert(pos + i, el)


def remove_tc_fields(paragraph):
    """Elimina los campos TC de un párrafo (runs entre begin y end inclusive). Devuelve cuántos."""
    p = paragraph._p
    n = 0
    for instr, _, _ in _campos_tc(p):
        primer_run = instr[0].getparent()
        # retroceder hasta el run con fldChar begin y avanzar hasta el end
        runs = list(p.iterchildren(qn("w:r")))
        i = runs.index(primer_run)
        while i > 0 and runs[i].find(qn("w:fldChar")) is None:
            i -= 1
        j = i
        while j < len(runs):
            fc = runs[j].find(qn("w:fldChar"))
            if fc is not None and fc.get(qn("w:fldCharType")) == "end":
                break
            j += 1
        for r in runs[i:j + 1]:
            p.remove(r)
        n += 1
    return n


# Elementos que en CT_Settings van DESPUÉS de w:updateFields (orden del esquema OOXML).
_DESPUES_DE_UPDATE_FIELDS = [
    "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids",
    "m:mathPr", "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping",
    "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade", "w:captions",
    "w:readModeInkLockDown", "w:smartTagType", "sl:schemaLibrary", "w:shapeDefaults",
    "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator",
]


def _insertar_update_fields(settings_xml):
    """Inserta <w:updateFields> respetando el orden del esquema, para que Word no pida reparar."""
    elemento = '<w:updateFields w:val="true"/>'
    posiciones = [settings_xml.find("<" + tag) for tag in _DESPUES_DE_UPDATE_FIELDS]
    posiciones = [p for p in posiciones if p >= 0]
    if posiciones:
        p = min(posiciones)
        return settings_xml[:p] + elemento + settings_xml[p:]
    return settings_xml.replace("</w:settings>", elemento + "</w:settings>", 1)


def set_update_fields_on_open(docx_path):
    """Agrega <w:updateFields w:val="true"/> a word/settings.xml, copiando el resto del paquete tal cual."""
    tmp_fd, tmp = tempfile.mkstemp(suffix=".docx")
    os.close(tmp_fd)
    with zipfile.ZipFile(docx_path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/settings.xml":
                s = data.decode("utf-8")
                if "w:updateFields" not in s:
                    s = _insertar_update_fields(s)
                data = s.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, docx_path)


# ---------------------------------------------------------------- fichas y edición

def find_ficha(doc, ficha_id):
    """Párrafo Heading 3 de la ficha (por ejemplo 'M1-11')."""
    for p in iter_parrafos(doc):
        if p.style is not None and p.style.name == "Heading 3":
            m = RE_FICHA.match(p.text.strip())
            if m and m.group(1) == ficha_id:
                return p
    return None


def fichas(doc):
    """[(id, titulo, fase_encabezado)] de todas las fichas, en orden."""
    out = []
    for p in iter_parrafos(doc):
        if p.style is not None and p.style.name == "Heading 3":
            t = p.text.strip()
            m = RE_FICHA.match(t)
            if m:
                resto = t[m.end():].strip()
                titulo, _, fase = resto.rpartition("–")
                out.append((m.group(1), titulo.strip(), fase.strip()))
    return out


def set_style(paragraph, style_name):
    paragraph.style = paragraph.part.document.styles[style_name]


def replace_in_runs(paragraph, old, new):
    """Reemplaza texto conservando el formato del primer run involucrado. Devuelve True si cambió."""
    runs = paragraph.runs
    full = "".join(r.text for r in runs)
    i = full.find(old)
    if i < 0:
        return False
    pos = 0
    first = None
    for r in runs:
        start, end = pos, pos + len(r.text)
        if end > i and start < i + len(old):
            a = max(i - start, 0)
            b = min(i + len(old) - start, len(r.text))
            if first is None:
                r.text = r.text[:a] + new + r.text[b:]
                first = r
            else:
                r.text = r.text[:a] + r.text[b:]
        pos = end
    return True


def move_block(start_el, end_el, before_el):
    """Mueve los elementos del cuerpo desde start_el hasta end_el (inclusive) antes de before_el."""
    bloque = []
    el = start_el
    while el is not None:
        bloque.append(el)
        if el is end_el:
            break
        el = el.getnext()
    for el in bloque:
        before_el.addprevious(el)


def insert_ficha_after(anchor_el, ficha_id, titulo, fase, secciones, plantilla_doc, modelos=None, num_id=None):
    """Inserta una ficha nueva después de anchor_el copiando el formato de una ficha existente.

    secciones: lista de (rotulo, [parrafos]) con rotulo en SITUACIÓN ACTUAL / REQUERIMIENTO /
    CRITERIOS DE ACEPTACIÓN / DEPENDENCIAS. Las viñetas de criterios usan 'List Paragraph'.
    Con `modelos` (ver modelos_ficha) los párrafos se clonan con el formato directo de la ficha
    modelo y los criterios se numeran con `num_id`. Devuelve el último elemento insertado.
    """
    if modelos is not None:
        num = fase.split()[-1]
        h = clonar_parrafo(modelos["h3"], [f"{ficha_id}     ", f"{titulo} ", "– FASE ", num])
        anchor_el.addnext(h)
        ultimo = h
        for rotulo, parrafos in secciones:
            nuevos = [clonar_parrafo(modelos["rotulo"], [rotulo])]
            for x in parrafos:
                if rotulo.startswith("CRITERIOS"):
                    nuevos.append(clonar_parrafo(modelos["lista"], [x], num_id=num_id))
                else:
                    nuevos.append(clonar_parrafo(modelos["cuerpo"], [x]))
            for el in nuevos:
                ultimo.addnext(el)
                ultimo = el
        return ultimo
    modelo = next(p for p in iter_parrafos(plantilla_doc) if p.style.name == "Heading 3")
    h = copy.deepcopy(modelo._p)
    for t in h.iter(qn("w:t")):
        t.text = ""
    primer_t = next(h.iter(qn("w:t")))
    primer_t.text = f"{ficha_id}     {titulo} – {fase}"
    anchor_el.addnext(h)
    ultimo = h
    styles = plantilla_doc.styles
    for rotulo, parrafos in secciones:
        for texto, estilo in [(rotulo, "Normal")] + [(x, "List Paragraph" if rotulo.startswith("CRITERIOS") else "Normal") for x in parrafos]:
            p = OxmlElement("w:p")
            ultimo.addnext(p)
            par = Paragraph(p, plantilla_doc)
            par.style = styles[estilo]
            par.add_run(texto)
            ultimo = p
    return ultimo


# ---------------------------------------------------------------- clonado con formato directo
# La especificación aplica el formato como formato directo (pPr/rPr) y no solo por estilo,
# por eso los párrafos nuevos se clonan desde un párrafo modelo del mismo tipo.

ESTILOS_TITULO = ("Heading 1", "Heading 2", "Heading 3")
ROTULOS = ("SITUACIÓN ACTUAL", "REQUERIMIENTO", "CRITERIOS DE ACEPTACIÓN", "DEPENDENCIAS")


W14 = "http://schemas.microsoft.com/office/word/2010/wordml"


def quitar_marcadores(el):
    """Elimina bookmarkStart/bookmarkEnd (los _Toc no pueden duplicarse) y los w14:paraId/textId
    copiados, que Word exige únicos; Word los regenera al guardar."""
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for b in list(el.iter(qn(tag))):
            b.getparent().remove(b)
    for x in el.iter():
        for attr in (f"{{{W14}}}paraId", f"{{{W14}}}textId"):
            if attr in x.attrib:
                del x.attrib[attr]
    return el


def _runs_texto(p_el):
    """Runs directos del párrafo que contienen w:t."""
    return [r for r in p_el.iterchildren(qn("w:r")) if r.find(qn("w:t")) is not None]


def fijar_textos_runs(p_el, textos):
    """Escribe textos[i] en el i-ésimo run con texto; elimina los runs de texto sobrantes."""
    runs = _runs_texto(p_el)
    if len(runs) < len(textos):
        raise ValueError(f"el párrafo tiene {len(runs)} runs de texto y se pidieron {len(textos)}")
    for r, texto in zip(runs, textos):
        ts = r.findall(qn("w:t"))
        ts[0].text = texto
        ts[0].set(qn("xml:space"), "preserve")
        for t in ts[1:]:
            r.remove(t)
    for r in runs[len(textos):]:
        p_el.remove(r)


def fijar_texto(paragraph, texto):
    """Reemplaza el texto completo de un párrafo conservando el formato de su primer run."""
    fijar_textos_runs(paragraph._p, [texto])


def fijar_tc(p_el, texto):
    """Reescribe el texto del (único) campo TC de un párrafo. Devuelve True si había campo."""
    campos = _campos_tc(p_el)
    for instr, _, nivel in campos:
        instr[0].text = f' TC "{texto}" \\f P \\l {nivel} '
        instr[0].set(qn("xml:space"), "preserve")
        for extra in instr[1:]:
            extra.text = ""
    return bool(campos)


def clonar_parrafo(modelo_el, textos, tc=None, num_id=None):
    """Copia profunda de un párrafo modelo con nuevos textos de run (y campo TC / numId opcionales)."""
    p = quitar_marcadores(copy.deepcopy(modelo_el))
    for c in list(p.iter(qn("w:commentRangeStart"), qn("w:commentRangeEnd"), qn("w:commentReference"))):
        c.getparent().remove(c)
    fijar_textos_runs(p, textos)
    if tc is not None:
        fijar_tc(p, tc)
    if num_id is not None:
        nid = p.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
        nid.set(qn("w:val"), str(num_id))
    return p


def nuevo_num_id(doc, abstract_id):
    """Agrega a numbering.xml una lista nueva (reinicia en 1) sobre abstractNum `abstract_id`."""
    numbering = doc.part.numbering_part.element
    ids = [int(n.get(qn("w:numId"))) for n in numbering.iterchildren(qn("w:num"))]
    nuevo = max(ids) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(nuevo))
    abs_ = OxmlElement("w:abstractNumId")
    abs_.set(qn("w:val"), str(abstract_id))
    num.append(abs_)
    over = OxmlElement("w:lvlOverride")
    over.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    over.append(start)
    num.append(over)
    numbering.append(num)
    return nuevo


def es_titulo(el):
    if el.tag != qn("w:p"):
        return False
    st = el.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
    return st is not None and st.get(qn("w:val")) in ("Heading1", "Heading2", "Heading3")


def tiene_dibujo(el):
    return el.tag == qn("w:p") and el.find(".//" + qn("w:drawing")) is not None


def texto_el(el):
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def bloque_ficha(doc, ficha_id):
    """(primer_el, ultimo_el) de una ficha: su Heading 3 y todo hasta antes del título siguiente."""
    h = find_ficha(doc, ficha_id)
    if h is None:
        raise KeyError(ficha_id)
    ultimo = h._p
    el = ultimo.getnext()
    while el is not None and el.tag in (qn("w:p"), qn("w:tbl")) and not es_titulo(el):
        ultimo = el
        el = el.getnext()
    return h._p, ultimo


def parrafos_seccion(doc, ficha_id, rotulo):
    """(rotulo_el, [párrafos de contenido]) de una sección de ficha.

    El contenido son los párrafos con texto que siguen al rótulo, hasta el siguiente rótulo,
    título, tabla, imagen o párrafo vacío.
    """
    inicio, fin = bloque_ficha(doc, ficha_id)
    el = inicio.getnext()
    rotulo_el = None
    while el is not None:
        if el.tag == qn("w:p") and texto_el(el).strip() == rotulo:
            rotulo_el = el
            break
        if el is fin:
            break
        el = el.getnext()
    if rotulo_el is None:
        raise KeyError(f"{ficha_id}: sin sección {rotulo}")
    contenido = []
    el = rotulo_el.getnext()
    while (el is not None and el.tag == qn("w:p") and not es_titulo(el) and not tiene_dibujo(el)
           and texto_el(el).strip() and texto_el(el).strip() not in ROTULOS):
        contenido.append(el)
        el = el.getnext()
    return rotulo_el, contenido


def reemplazar_seccion(doc, ficha_id, rotulo, textos):
    """Reemplaza el contenido de una sección de ficha clonando el formato de su primer párrafo."""
    rotulo_el, contenido = parrafos_seccion(doc, ficha_id, rotulo)
    modelo = copy.deepcopy(contenido[0])
    for el in contenido:
        el.getparent().remove(el)
    ultimo = rotulo_el
    for texto in textos:
        nuevo = clonar_parrafo(modelo, [texto])
        ultimo.addnext(nuevo)
        ultimo = nuevo
    return ultimo


def fijar_texto_celda(tc_el, texto):
    """Deja una celda con un solo párrafo y un solo run con `texto`, conservando el formato del primero."""
    ps = tc_el.findall(qn("w:p"))
    con_texto = [p for p in ps if _runs_texto(p)] or ps
    p = con_texto[0]
    for otro in ps:
        if otro is not p:
            tc_el.remove(otro)
    if _runs_texto(p):
        fijar_textos_runs(p, [texto])
    else:
        Paragraph(p, None).add_run(texto)


def agregar_fila(tbl_el, modelo_tr, textos, antes_de=None):
    """Agrega una fila clonada de `modelo_tr` con los textos dados (al final o antes de `antes_de`)."""
    tr = quitar_marcadores(copy.deepcopy(modelo_tr))
    for tc, texto in zip(tr.findall(qn("w:tc")), textos):
        fijar_texto_celda(tc, texto)
    if antes_de is not None:
        antes_de.addprevious(tr)
    else:
        tbl_el.append(tr)
    return tr


def texto_documento(doc):
    """Todo el texto visible del cuerpo (incluye tablas e índice en caché), encabezados y pies."""
    partes = ["".join(t.text or "" for t in p.iter(qn("w:t"))) for p in doc.element.body.iter(qn("w:p"))]
    for s in doc.sections:
        for hf in (s.header, s.footer, s.first_page_header, s.first_page_footer, s.even_page_header, s.even_page_footer):
            partes += [p.text for p in hf.paragraphs]
    return "\n".join(partes)
